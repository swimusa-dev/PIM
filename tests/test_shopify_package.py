"""EXP-13 Shopify update package: diff, protected fields (D-09), review workbook, import CSV rules, reconcile."""

import csv
import io

from openpyxl import load_workbook
from sqlalchemy import select
from tests.conftest import FIXTURES

from pim.adapters.shopify_admin import parse_product_csv
from pim.export.shopify_package import apply_review_workbook, build_package, reconcile
from pim.models import ProductEntity, ReferenceEntity, UpdatePackage
from pim.services.readiness import refresh_snapshots
from pim.services.storage import Storage
from pim.services.values import set_value


def _sc(session, style, color):
    return session.execute(select(ProductEntity).where(ProductEntity.level == "style_color", ProductEntity.style_number == style,
                                                       ProductEntity.color_code == color)).scalar_one()


def _make_ready(session, meta, admin, sc, title, color_name, shots=True):
    st = session.get(ProductEntity, sc.parent_id)
    set_value(session, meta, admin, st, meta.attributes["product_description"], "<p>New description from the PIM.</p>")
    fabric = session.execute(select(ReferenceEntity).where(ReferenceEntity.code == "ecorib")).scalar_one_or_none()
    if fabric is None:
        fabric = ReferenceEntity(type_code="fabric", code="ecorib", label="EcoRib", certificate_refs=["RCS-2026-001"])
        session.add(fabric)
        session.flush()
        pseudo = ProductEntity(id=fabric.id, level="ref:fabric", style_number="ecorib", label="EcoRib")
        set_value(session, meta, admin, pseudo, meta.attributes["fabric_story"], "EcoRib is our signature ribbed fabric.")
        set_value(session, meta, admin, pseudo, meta.attributes["fabric_care"], "hand_wash_cold_dry_flat")
    for code, val in {"product_title": title, "customer_color_name": color_name, "fabric": fabric.id,
                      "seo_title": f"{title} | Vitamin A", "seo_description": "Old meta description"}.items():
        assert set_value(session, meta, admin, sc, meta.attributes[code], val).outcome == "approved", code
    if shots:
        from pim.models import Asset, AssetLink

        for i, shot in enumerate(["front", "back", "on_model_front"]):
            a = Asset(id=f"a-{sc.id[:6]}-{i}", filename=f"{sc.style_number}-{sc.color_code}_{shot}.jpg")
            session.add(a)
            session.flush()
            session.add(AssetLink(asset_id=a.id, entity_id=sc.id, shot_type=shot, tier=1, status="confirmed"))
    session.flush()
    refresh_snapshots(session, meta, [sc])


def _build(session, meta, admin, tmp_path, selection=None):
    snap = parse_product_csv((FIXTURES / "shopify_products_export.csv").read_bytes())
    pkg = UpdatePackage(channel_code="shopify-vitamina", created_by=admin.email, snapshot_mode="csv_upload",
                        selection=selection or {"entity_ids": [e.id for e in session.execute(select(ProductEntity).where(ProductEntity.level == "style_color")).scalars()]})
    session.add(pkg)
    session.flush()
    storage = Storage(str(tmp_path))
    build_package(session, meta, admin, pkg, snap, storage)
    session.flush()
    return pkg, snap, storage


def test_snapshot_csv_parser_reads_products_variants_and_metafields():
    snap = parse_product_csv((FIXTURES / "shopify_products_export.csv").read_bytes())
    p = snap.by_handle("skylar-top-festival-pink")
    assert p and len(p["variants"]) == 3 and p["variants"][1]["sku"] == "200809-FVP-S"
    assert p["metafields"]["custom.care_instructions"] == "Hand wash cold. Line dry."
    assert "shopify--discovery--product_recommendation.complementary_products" in p["metafields"]
    assert snap.by_sku("200810-ECB-ML")[0]["handle"] == "luciana-bottom-eco-black"


def test_exp13_diff_skips_not_ready_and_protects_layer_fields(session, meta, admin, ingested, tmp_path):
    fvp = _sc(session, "200809", "FVP")
    _make_ready(session, meta, admin, fvp, "Skylar Top", "Festival Pink")
    pkg, snap, storage = _build(session, meta, admin, tmp_path)
    assert pkg.status == "ready" and pkg.product_count == 1
    fields = {r.field for r in pkg.rows}
    assert "Body (HTML)" in fields and "Fabric story (product.metafields.custom.fabric_story)" in fields
    assert not any("product_recommendation" in f for f in fields)  # D-09 protected
    assert not any(r.field == "SEO Description" for r in pkg.rows)  # unchanged value produces no row
    skipped = pkg.summary["skipped"]
    assert any("Not ready" in s["reason"] for s in skipped) and any("Not eligible" in s["reason"] for s in skipped)
    assert any(f["name"] == "review_workbook.xlsx" for f in pkg.files)
    care = next(r for r in pkg.rows if "care_instructions" in r.field)
    assert care.current_value == "Hand wash cold. Line dry." and care.proposed_value == "Hand wash cold. Dry flat."
    # variant option value change (customer color name -> Option1 Value) lands in a variant file with option columns
    variant_rows = [r for r in pkg.rows if r.scope == "variant"]
    assert not variant_rows  # "Festival Pink" already matches Shopify


def test_exp13_import_csv_rules(session, meta, admin, ingested, tmp_path):
    fvp = _sc(session, "200809", "FVP")
    _make_ready(session, meta, admin, fvp, "Skylar Longline Top", "Festival Pink Print")
    pkg, snap, storage = _build(session, meta, admin, tmp_path)
    product_files = [f for f in pkg.files if f["name"].startswith("products_")]
    variant_files = [f for f in pkg.files if f["name"].startswith("variants_")]
    assert product_files and variant_files
    rows = list(csv.DictReader(io.StringIO(storage.read(product_files[0]["path"]).decode())))
    assert rows[0]["Handle"] == "skylar-top-festival-pink" and rows[0]["Title"] == "Skylar Longline Top"
    # no blank cells unless an intentional clear, and only changed columns are present
    for r in rows:
        for k, v in r.items():
            assert v != "" or any(c["field"] == k for c in pkg.summary["intentional_clears"]), k
    assert "Variant SKU" not in rows[0] and "Option1 Name" not in rows[0]
    vrows = list(csv.DictReader(io.StringIO(storage.read(variant_files[0]["path"]).decode())))
    assert {"Handle", "Title", "Option1 Name", "Option1 Value", "Option2 Name", "Option2 Value", "Variant SKU"} <= set(vrows[0])
    assert all(v["Option1 Name"] == "Color" and v["Option1 Value"] == "Festival Pink Print" for v in vrows)
    assert len(vrows) == 3


def test_exp13_review_roundtrip_and_reconcile(session, meta, admin, ingested, tmp_path):
    fvp = _sc(session, "200809", "FVP")
    _make_ready(session, meta, admin, fvp, "Skylar Top", "Festival Pink")
    pkg, snap, storage = _build(session, meta, admin, tmp_path)
    wb_bytes = storage.read(next(f["path"] for f in pkg.files if f["kind"] == "workbook"))
    wb = load_workbook(io.BytesIO(wb_bytes))
    ws = wb["Changes"]
    header = [c.value for c in ws[1]]
    inc = header.index("Include (Y/N)") + 1
    ws.cell(row=2, column=inc, value="N")
    dropped_field = ws.cell(row=2, column=header.index("Field") + 1).value
    buf = io.BytesIO()
    wb.save(buf)
    changed = apply_review_workbook(session, pkg, buf.getvalue(), snap, storage)
    assert changed == 1 and pkg.version == 2
    assert not any(r.field == dropped_field and r.include for r in pkg.rows)
    # reconcile: simulate Shopify applying the Body change, leaving the rest
    applied = snap
    applied.by_handle("skylar-top-festival-pink")["body_html"] = "<p>New description from the PIM.</p>"
    applied.by_handle("skylar-top-festival-pink")["metafields"]["custom.care_instructions"] = "Someone else's text"
    counts = reconcile(pkg, applied)
    by_field = {r.field: r.reconcile_status for r in pkg.rows}
    assert by_field["Body (HTML)"] == "applied"
    assert by_field["Care instructions (product.metafields.custom.care_instructions)"] == "changed_differently"
    assert counts.get("not_applied", 0) >= 1 and pkg.status == "reconciled"

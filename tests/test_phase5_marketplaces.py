"""Phase 5: EXP-08 Walmart Marketplace, EXP-09 retailer feeds with delivery, EXP-10 and EXP-11 through Rithum (D-01, D-02),
D-13 Matrixify workbook, D-03 Shopify sandbox write path, D-05 merch lead cost visibility."""

from __future__ import annotations

import io
import json
from datetime import timedelta

import httpx
import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from sqlalchemy import select
from tests.conftest import FIXTURES, make_user

from pim.adapters.base import NotConfigured
from pim.adapters.rithum import DEFAULT_COLUMNS, RithumFeedAdapter, image_name
from pim.adapters.shopify_admin import ShopifyAdminClient, parse_product_csv
from pim.adapters.walmart_mp import WalmartMarketplaceAdapter, next_poll_at
from pim.config import get_settings
from pim.export import publishing
from pim.export.payloads import build_payloads, payload_files
from pim.export.schema_loaders import load_walmart_spec
from pim.export.shopify_package import build_matrixify_workbook
from pim.models import Asset, AssetLink, Channel, ProductEntity, PublishItem, UserRole
from pim.models._common import utcnow
from pim.services import delivery
from pim.services.values import set_value

ADMIN = {"X-Forwarded-Email": "admin@swimusa.com", "X-Forwarded-User": "u1"}
VIEWER = {"X-Forwarded-Email": "viewer@swimusa.com", "X-Forwarded-User": "u2"}
CHANNEL_STATE = ("target_schema", "attribute_mapping", "transforms", "image_profile", "config", "enabled", "version")


@pytest.fixture(autouse=True)
def _restore_channels(session):
    import copy

    saved = {c.code: {k: copy.deepcopy(getattr(c, k)) for k in CHANNEL_STATE} for c in session.execute(select(Channel)).scalars()}
    yield
    session.rollback()
    for c in session.execute(select(Channel)).scalars():
        for k, v in saved.get(c.code, {}).items():
            setattr(c, k, v)
    session.commit()


@pytest.fixture
def client(session, ingested):
    from pim.main import app

    for email, role in (("admin@swimusa.com", "pim_admin"), ("viewer@swimusa.com", "viewer")):
        if not session.execute(select(UserRole).where(UserRole.email == email, UserRole.role_code == role)).first():
            session.add(UserRole(email=email, role_code=role, brand_scopes=[]))
    session.commit()
    return TestClient(app)


def _sc(session, style="200809", color="FVP"):
    return session.execute(select(ProductEntity).where(ProductEntity.level == "style_color", ProductEntity.style_number == style,
                                                       ProductEntity.color_code == color)).scalar_one()


def _ready(session, meta, admin, sc, title="Skylar Top"):
    st = session.get(ProductEntity, sc.parent_id)
    set_value(session, meta, admin, sc, meta.attributes["product_title"], title)
    set_value(session, meta, admin, sc, meta.attributes["customer_color_name"], "Festival Pink")
    set_value(session, meta, admin, st, meta.attributes["product_description"], "<p>A longline bralette in soft ribbed fabric.</p>")
    a = Asset(id=f"img-{sc.id[:6]}", filename=f"{sc.style_number}-{sc.color_code}_front_01.jpg", width=2400, height=2400,
              thumbnail_url=f"https://cdn.example/{sc.id[:6]}.jpg")
    session.add(a)
    session.flush()
    session.add(AssetLink(asset_id=a.id, entity_id=sc.id, shot_type="front", status="confirmed"))
    session.flush()


# ---------------------------------------------------------------------------------------------------
# EXP-08 Walmart
# ---------------------------------------------------------------------------------------------------
class WalmartMock:
    def __init__(self) -> None:
        self.feeds: list[dict] = []
        self.calls: list[str] = []
        self.item_state = "INPROGRESS"

    def handler(self, request: httpx.Request) -> httpx.Response:
        path = request.url.path
        self.calls.append(f"{request.method} {path}")
        if path == "/v3/token":
            assert request.headers["Authorization"].startswith("Basic ") and request.headers.get("WM_SVC.NAME")
            return httpx.Response(200, json={"access_token": "tok", "expires_in": 900})
        assert request.headers.get("WM_SEC.ACCESS_TOKEN") == "tok" and request.headers.get("WM_QOS.CORRELATION_ID")
        if path == "/v3/feeds" and request.method == "POST":
            assert request.url.params["feedType"] == "MP_ITEM"
            body = request.content.decode()
            self.feeds.append(json.loads(body[body.index("{"): body.rindex("}") + 1]))
            return httpx.Response(200, json={"feedId": f"FEED-{len(self.feeds)}"})
        if path.startswith("/v3/feeds/"):
            return httpx.Response(200, json={"feedStatus": "PROCESSED", "itemsReceived": 2, "itemsSucceeded": 1, "itemsFailed": 0,
                                             "itemDetails": {"itemIngestionStatus": [
                                                 {"sku": "200809-FVP", "ingestionStatus": "SUCCESS", "ingestionErrors": {"ingestionError": []}},
                                                 {"sku": "200809-ECB", "ingestionStatus": self.item_state,
                                                  "ingestionErrors": {"ingestionError": [{"code": "ERR_PRICE", "description": "price missing", "field": "price"}] if self.item_state == "DATA_ERROR" else []}}]}})
        if "listingQuality" in path:
            return httpx.Response(200, json={"payload": [{"sku": "200809-FVP", "itemQualityScore": {"score": 82}}, {"sku": "200809-ECB", "qualityScore": 64}]})
        return httpx.Response(404)


def _walmart(mock: WalmartMock, max_items: int = 5) -> WalmartMarketplaceAdapter:
    return WalmartMarketplaceAdapter(client_id="id", client_secret="secret", spec_version="5.0", max_items_per_feed=max_items,
                                     transport=httpx.MockTransport(mock.handler))


def test_exp08_walmart_feeds_token_batching_status_and_cadence(monkeypatch):
    mock = WalmartMock()
    w = _walmart(mock, max_items=5)
    items = [{"orderable": {"sku": f"S{i}", "productIdentifiers": {"productIdType": "UPC", "productId": f"8400000000{i:02d}"}, "price": 120},
              "visible": {"Swimwear": {"productName": f"Top {i}", "mainImageUrl": "https://cdn.example/a.jpg", "swatchImages": [{"swatchImageUrl": "x"}]}},
              "variant_group": "200809" if i < 3 else f"G{i}"} for i in range(7)]
    w.group_variants(items, ["color", "clothingSize"])
    grp = [it["orderable"] for it in items if it["orderable"].get("variantGroupId") == "200809"]
    assert len(grp) == 3 and sum(1 for o in grp if o["isPrimaryVariant"] == "Yes") == 1 and all(o["variantAttributeNames"] == ["color", "clothingSize"] for o in grp)
    assert w.validate(items).ok
    items[1]["orderable"]["isPrimaryVariant"] = "Yes"
    assert "exactly one isPrimaryVariant" in w.validate(items).errors[0]["message"]
    items[1]["orderable"]["isPrimaryVariant"] = "No"
    feeds = w.build_feed(items)
    assert len(feeds) == 2 and feeds[0]["MPItemFeedHeader"]["version"] == "5.0" and len(feeds[0]["MPItem"]) == 5 and "Orderable" in feeds[0]["MPItem"][0]
    with pytest.raises(NotConfigured):
        w.publish(items)  # PIM_FEATURE_CHANNEL_SEND off
    assert w.publish(items, dry_run=True).detail["feeds"] == 2
    monkeypatch.setattr(get_settings(), "feature_channel_send", True)
    res = w.publish(items)
    assert res.status == "submitted" and res.detail["feed_ids"] == ["FEED-1", "FEED-2"] and res.batch_id == "FEED-1"
    assert mock.calls.count("POST /v3/token") == 1  # token cached across the two feed posts
    st = w.fetch_status("FEED-1")
    assert st["status"] == "PROCESSED" and st["complete"] is False  # one item still INPROGRESS [C14]
    assert {i["external_id"]: i["status"] for i in st["items"]} == {"200809-FVP": "accepted", "200809-ECB": "in_review"}
    mock.item_state = "DATA_ERROR"
    errs = w.fetch_errors("FEED-1")
    assert errs[0]["external_id"] == "200809-ECB" and "price missing" in errs[0]["message"] and errs[0]["field"] == "price"
    assert w.listing_quality(["200809-FVP", "200809-ECB"]) == {"200809-FVP": 82.0, "200809-ECB": 64.0}
    now = utcnow()
    assert [next_poll_at(n, now) - now for n in (0, 1, 2, 3, 9)] == [timedelta(minutes=15), timedelta(hours=1), timedelta(hours=2), timedelta(hours=4), timedelta(hours=4)]
    maint = w.build_maintenance_feed(items[:1])
    assert maint["MPItemFeedHeader"]["processMode"] == "PARTIAL_UPDATE"


def test_exp08_publish_flow_submits_polls_and_scores(session, meta, admin, ingested, monkeypatch):
    ch = session.get(Channel, "walmart-mp")
    ch.target_schema = load_walmart_spec((FIXTURES / "phase3" / "walmart_spec_v1.json").read_bytes(), "Swimwear").as_dict()
    ch.attribute_mapping = {"product_title": "productName", "product_description": "shortDescription", "customer_color_name": "color",
                            "core.style_number": "productId", "core.color_code": "sku", "erp_brand_name": "brand"}
    ch.config = {"transport": "api", "file_format": "json", "walmart": {"spec_version": "5.0", "max_items_per_feed": 5000,
                                                                        "variant_attribute_names": ["color"]}}
    ch.enabled = True
    fvp, ecb = _sc(session), _sc(session, "200809", "ECB")
    for sc in (fvp, ecb):
        _ready(session, meta, admin, sc)
    session.flush()
    mock = WalmartMock()
    adapter = _walmart(mock, max_items=5000)
    monkeypatch.setattr(publishing, "adapter_for", lambda channel: adapter)
    monkeypatch.setattr(get_settings(), "feature_channel_send", True)
    # make every product valid for the test by relaxing the raw spec requirements
    payloads = build_payloads(session, meta, ch, [fvp, ecb])
    for p in payloads:
        p.issues = []
    monkeypatch.setattr(publishing, "build_payloads", lambda *a, **k: payloads)
    rec = publishing.start_run(session, ch, "publish", "full", admin.email, {"entity_ids": [fvp.id, ecb.id]})
    publishing.execute_run(session, meta, ch, rec, send=True)
    session.commit()
    assert rec.status == "submitted" and rec.external_batch_id == "FEED-1", rec.error
    assert rec.summary["walmart"]["polls"] == 0 and rec.summary["walmart"]["next_poll_at"] and rec.summary["delivery"]["transport"] == "api"
    assert publishing.walmart_poll_due(rec) is False  # 15 minutes out
    assert publishing.poll_feedback(session, meta, ch, rec, force=False)["skipped"] == "not due"
    result = publishing.poll_feedback(session, meta, ch, rec)
    assert result["accepted"] == 1 and result["in_review"] == 1
    assert rec.summary["walmart"]["polls"] == 1 and rec.summary["walmart"]["status"] == "in_progress" and rec.summary["listing_quality"]["200809-FVP"] == 82.0
    items = {i.external_id: i.status for i in session.execute(select(PublishItem).where(PublishItem.publish_id == rec.id)).scalars()}
    assert items == {"200809-FVP": "accepted", "200809-ECB": "in_review"}
    mock.item_state = "SUCCESS"
    publishing.poll_feedback(session, meta, ch, rec)
    assert rec.summary["walmart"]["status"] == "complete" and rec.summary["walmart"]["next_poll_at"] is None and rec.status == "accepted"
    h = publishing.channel_health(session, ch)
    assert h["listing_quality_avg"] == 73.0 and h["feeds_in_progress"] == 0


# ---------------------------------------------------------------------------------------------------
# EXP-09, EXP-10, EXP-11 Rithum feeds
# ---------------------------------------------------------------------------------------------------
def _template(headers: list[str]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Nordstrom"
    ws.append(["Supplier catalog template v3"])
    ws.append(headers)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_exp10_rithum_feed_one_row_per_sku_with_template_nrf_and_image_names(session, meta, admin, ingested):
    ch = session.get(Channel, "nordstrom")
    fvp = _sc(session)
    _ready(session, meta, admin, fvp)
    # NRF codes ride on option value channel mappings (GS1 NRF, [C18][C19])
    from pim.models import OptionValue

    sz = session.execute(select(OptionValue).where(OptionValue.code == "XS")).scalars().first()
    sz.channel_mappings = {**(sz.channel_mappings or {}), "nrf": "0030"}
    session.flush()
    from pim.services.metadata import load_metadata

    meta = load_metadata(session, force=True)
    headers = [c["header"] for c in DEFAULT_COLUMNS["nordstrom_dropship"]]
    ch.config = {**ch.config, "header_row": 2}
    payloads = build_payloads(session, meta, ch, [fvp])
    p = payloads[0]
    assert len(p.variants) >= 2 and all(v["sku"] for v in p.variants)
    assert p.feed["Product Name"] == "Skylar Top" and p.feed["Supplier Color Name"] == "Festival Pink" and p.feed["Image 1"].startswith("https://cdn.example/")
    assert p.images[0]["name"] == "200809-FestivalPink-01"  # VPN-SupplierColorName-01 [C21]
    assert image_name("200809", "festival pink", 2, prefix="VA") == "VA200809-FestivalPink-02"
    files = {n: c for n, c, _k in payload_files(ch, payloads, None, _template(headers))}
    assert set(files) == {"feed.xlsx", "validation_report.csv"}
    wb = load_workbook(io.BytesIO(files["feed.xlsx"]))
    ws = wb.active
    assert ws.title == "Nordstrom" and ws.cell(1, 1).value == "Supplier catalog template v3"  # the template is preserved
    rows = list(ws.iter_rows(min_row=3, values_only=True))
    assert len(rows) == len(p.variants)
    col = {h: i for i, h in enumerate(headers)}
    skus = {r[col["Supplier SKU"]] for r in rows}
    assert skus == {v["sku"] for v in p.variants} and all(r[col["UPC"]] for r in rows) and all(r[col["Vendor Part Number"]] == "200809" for r in rows)
    xs = next(r for r in rows if r[col["Size"]] == "XS")
    assert xs[col["NRF Size Code"]] == "0030" and xs[col["AI Generated Imagery"]] == "No"


def test_exp11_targetplus_family_rules_fail_the_whole_family(session, meta, admin, ingested):
    ch = session.get(Channel, "targetplus")
    fvp = _sc(session)
    _ready(session, meta, admin, fvp, title="A " * 80)  # 160 characters, over the Target+ limit (UNCONFIRMED [C17])
    payloads = build_payloads(session, meta, ch, [fvp])
    p = payloads[0]
    fam = [i for i in p.issues if "whole family fails" in i.message]
    assert fam and "title longer than 150" in fam[0].message
    adapter = RithumFeedAdapter(program="targetplus")
    assert adapter._family_errors([{"core.style_number": "1", "core.sku": "1-S", "product_title": "Fine", "gtin": "840000000011", "image_width": 1200, "image_height": 1200}]) == []
    bad = adapter._family_errors([{"core.style_number": "1", "core.sku": "1-S", "product_title": "Fine", "gtin": "ABC", "image_width": 1200, "image_height": 900}])
    assert bad and "GTIN" in bad[0]["details"][0] and any("1:1" in d for d in bad[0]["details"])
    assert RithumFeedAdapter(program="nordstrom_dropship").validate([{"core.sku": "x", "product_title": "t" * 200}]).ok is False  # required columns only


def test_exp09_delivery_sftp_email_and_download(monkeypatch):
    class Chan:
        label = "Nordstrom"
        config = {"transport": "sftp", "sftp": {"host": "sftp.example", "port": 22, "path": "/in", "username": "swim", "secret_name": "nordstrom-sftp"}}

    class FakeSftp:
        written: dict[str, bytes] = {}

        def open(self, path, mode):  # noqa: ANN001
            buf = io.BytesIO()
            buf.close = lambda: FakeSftp.written.__setitem__(path, buf.getvalue())  # type: ignore[assignment]
            return buf

    r = delivery.deliver(Chan(), [("feed.xlsx", b"abc")], sftp_client=FakeSftp())
    assert r["status"] == "delivered" and FakeSftp.written["/in/feed.xlsx"] == b"abc" and r["files"] == ["feed.xlsx"]
    r = delivery.deliver(Chan(), [("feed.xlsx", b"abc")])  # no secret bound
    assert r["status"] == "failed" and "nordstrom-sftp" in r["detail"]
    monkeypatch.setenv("NORDSTROM_SFTP", "pw")
    assert delivery.secret_value("nordstrom-sftp") == "pw"
    sent = {}

    class FakeSmtp:
        def __enter__(self):
            return self

        def __exit__(self, *a):  # noqa: ANN002
            return False

        def send_message(self, msg):  # noqa: ANN001
            sent["to"], sent["subject"], sent["attachments"] = msg["To"], msg["Subject"], [p.get_filename() for p in msg.iter_attachments()]

    Chan.config = {"transport": "email", "email": {"to": ["buyer@retailer.com"], "subject": "Feed"}}
    r = delivery.deliver(Chan(), [("feed.csv", b"a,b")], smtp_factory=FakeSmtp)
    assert r["status"] == "delivered" and sent == {"to": "buyer@retailer.com", "subject": "Feed", "attachments": ["feed.csv"]}
    Chan.config = {"transport": "download"}
    assert delivery.deliver(Chan(), [("feed.csv", b"a")])["status"] == "skipped"


# ---------------------------------------------------------------------------------------------------
# D-13 Matrixify and D-03 sandbox
# ---------------------------------------------------------------------------------------------------
class Row:
    def __init__(self, handle, scope, field, value, is_clear=False, variant_key=None, include=True):  # noqa: ANN001
        self.handle, self.scope, self.field, self.proposed_value, self.is_clear, self.variant_key, self.include = handle, scope, field, value, is_clear, variant_key, include


def test_d13_matrixify_workbook():
    snap = parse_product_csv((FIXTURES / "shopify_products_export.csv").read_bytes())
    handle = next(iter(snap.products))
    rows = [Row(handle, "product", "Title", "New Title"), Row(handle, "product", "SEO Description", "", is_clear=True),
            Row(handle, "metafield", "custom.short_description", "Short"), Row(handle, "variant", "Variant Barcode", "840000000011", variant_key={"sku": "SKU-1"}),
            Row(handle, "product", "Tags", "ignored", include=False)]
    wb = load_workbook(io.BytesIO(build_matrixify_workbook(rows, snap)))
    ws = wb["Products"]
    header = [c.value for c in ws[1]]
    assert header[:3] == ["ID", "Handle", "Command"] and "Metafield: custom.short_description [single_line_text_field]" in header and "Tags" not in header
    row = [c.value for c in ws[2]]
    assert row[header.index("Handle")] == handle and row[header.index("Command")] == "UPDATE" and row[header.index("Title")] == "New Title"
    assert row[header.index("SEO Description")] in ("", None)
    wv = wb["Variants"]
    assert [c.value for c in wv[1]] == ["Handle", "Command", "Variant SKU", "Variant Barcode"] and wv.cell(2, 4).value == "840000000011"


def test_d03_sandbox_plan_and_protected_fields_unchanged(monkeypatch):
    protected = ["metafield:shopify--discovery--product_recommendation.complementary_products", "Variant Price"]
    rows = [{"handle": "skylar-top", "scope": "product", "field": "Title", "proposed_value": "Skylar Top"},
            {"handle": "skylar-top", "scope": "product", "field": "SEO Title", "proposed_value": "Skylar | Vitamin A"},
            {"handle": "skylar-top", "scope": "metafield", "field": "shopify--discovery--product_recommendation.complementary_products", "proposed_value": "[]"},
            {"handle": "skylar-top", "scope": "variant", "field": "Variant Price", "proposed_value": "1", "variant_key": {"sku": "S1"}},
            {"handle": "skylar-top", "scope": "variant", "field": "Variant Barcode", "proposed_value": "840000000011", "variant_key": {"sku": "S1"}}]
    rows += [{"handle": "skylar-top", "scope": "metafield", "field": f"custom.k{i}", "proposed_value": str(i)} for i in range(30)]
    plan = ShopifyAdminClient.plan_mutations(rows, protected)
    kinds = [m["kind"] for m in plan]
    assert kinds == ["productUpdate", "metafieldsSet", "metafieldsSet", "productVariantsBulkUpdate"]
    assert plan[0]["input"] == {"title": "Skylar Top", "seo.title": "Skylar | Vitamin A"}
    assert len(plan[1]["metafields"]) == 25 and len(plan[2]["metafields"]) == 5  # batches of 25 [C4]
    assert not any("complementary_products" in f for m in plan for f in m["fields"]) and plan[3]["fields"] == ["Variant Barcode"]
    # execute against a fake sandbox store; the protected metafield must read back unchanged (EXP-07 required test)
    state = {"complementary": '["gid://shopify/Product/2"]', "title": "Old"}
    calls = []

    def handler(request: httpx.Request) -> httpx.Response:
        body = json.loads(request.content)
        q = body["query"]
        calls.append(q.split("(")[0].strip())
        if "productByHandle" in q:
            return httpx.Response(200, json={"data": {"productByHandle": {"id": "gid://shopify/Product/1", "handle": "skylar-top", "title": state["title"],
                "metafields": {"edges": [{"node": {"id": "m1", "namespace": "shopify--discovery--product_recommendation", "key": "complementary_products", "value": state["complementary"], "type": "list.product_reference"}}]},
                "variants": {"edges": [{"node": {"id": "gid://shopify/ProductVariant/9", "sku": "S1"}}]}}}})
        if "productUpdate" in q:
            state["title"] = body["variables"]["input"]["title"]
            return httpx.Response(200, json={"data": {"productUpdate": {"product": {"id": "gid://shopify/Product/1"}, "userErrors": []}}})
        if "metafieldsSet" in q:
            assert all("complementary" not in m["key"] for m in body["variables"]["m"])
            return httpx.Response(200, json={"data": {"metafieldsSet": {"metafields": [], "userErrors": []}}})
        if "productVariantsBulkUpdate" in q:
            assert body["variables"]["v"] == [{"id": "gid://shopify/ProductVariant/9", "barcode": "840000000011"}]
            return httpx.Response(200, json={"data": {"productVariantsBulkUpdate": {"userErrors": []}}})
        return httpx.Response(400)

    with pytest.raises(NotConfigured):
        ShopifyAdminClient.sandbox_client()
    monkeypatch.setattr(get_settings(), "shopify_sandbox_domain", "dev-store.myshopify.com")
    monkeypatch.setattr(get_settings(), "shopify_sandbox_token", "shpat_dev")
    c = ShopifyAdminClient.sandbox_client(transport=httpx.MockTransport(handler))
    before = c.read_protected("skylar-top", protected)
    done = c.execute_plan(plan)
    after = c.read_protected("skylar-top", protected)
    assert all(m["status"] == "ok" for m in done) and state["title"] == "Skylar Top"
    assert before == after == {"shopify--discovery--product_recommendation.complementary_products": '["gid://shopify/Product/2"]'}
    prod = ShopifyAdminClient(store_domain="vitamin-a.myshopify.com", token="t", transport=httpx.MockTransport(handler))
    with pytest.raises(NotConfigured):
        prod.execute_plan(plan)  # production store: flag off, never written


# ---------------------------------------------------------------------------------------------------
# API and roles
# ---------------------------------------------------------------------------------------------------
def test_phase5_api_transport_template_deliver_sandbox_and_merch_lead(client, session, meta, admin):
    body = {"transport": "sftp", "file_format": "xlsx", "header_row": 1, "row_level": "sku",
            "feed_columns": [{"header": "SKU", "source": "core.sku", "required": True}, {"header": "Color", "source": "customer_color_name", "case_transform": "title"}],
            "sftp": {"host": "sftp.example", "port": 22, "path": "/in", "username": "swim", "secret_name": "nordstrom-sftp"},
            "rithum": {"program": "nordstrom_dropship", "family_rules": False}}
    assert client.put("/api/channels/nordstrom/transport", json=body, headers=VIEWER).status_code == 403
    r = client.put("/api/channels/nordstrom/transport", json=body, headers=ADMIN)
    assert r.status_code == 200, r.text
    assert r.json()["transport"] == "sftp" and r.json()["feed_columns"][0]["required"] is True and r.json()["row_level"] == "sku"
    assert client.put("/api/channels/nordstrom/transport", json={**body, "feed_columns": [{"header": "X", "source": "ghost"}]}, headers=ADMIN).status_code == 400
    assert client.put("/api/channels/nordstrom/transport", json={**body, "sftp": {"host": ""}}, headers=ADMIN).status_code == 400
    assert client.put("/api/channels/nordstrom/transport", json={**body, "transport": "email", "email": {"to": []}}, headers=ADMIN).status_code == 400
    csv_t = b"Supplier SKU,UPC,Product Name,Size\n"
    up = client.post("/api/channels/nordstrom/feed-template", headers=ADMIN, files={"file": ("nordstrom.csv", csv_t)})
    assert up.status_code == 200, up.text
    assert up.json()["columns"] == ["Supplier SKU", "UPC", "Product Name", "Size"]
    t = client.get("/api/channels/nordstrom/transport", headers=VIEWER).json()
    assert t["feed_template_name"] == "nordstrom.csv" and {c["header"] for c in t["feed_columns"]} >= {"SKU", "Color", "Supplier SKU", "UPC"}
    assert client.get("/api/channels/nordstrom/feed-template", headers=VIEWER).content == csv_t
    xl = _template(["Vendor Part Number", "Supplier SKU"])
    up2 = client.post("/api/channels/targetplus/feed-template", headers=ADMIN, files={"file": ("tp.xlsx", xl)})
    assert up2.status_code == 200 and up2.json()["schema_loaded"] is True
    # deliver: download channels refuse; a dry run refuses
    fvp = _sc(session)
    _ready(session, meta, admin, fvp)
    session.commit()
    dr = client.post("/api/channels/targetplus/dry-run", json={"selection": {"entity_ids": [fvp.id]}, "mode": "full"}, headers=ADMIN)
    assert dr.status_code == 200, dr.text
    assert client.post(f"/api/publishes/{dr.json()['id']}/deliver", headers=ADMIN).status_code == 409
    # sandbox plan through a package
    pkg = client.post("/api/packages", json={"channel_code": "shopify-vitamina", "selection": {"entity_ids": [fvp.id]}, "snapshot": {"mode": "csv_upload"}}, headers=ADMIN).json()
    plan = client.post(f"/api/packages/{pkg['id']}/sandbox-run", json={"plan_only": True}, headers=ADMIN)
    assert plan.status_code == 200 and plan.json()["mode"] == "plan"
    assert client.post(f"/api/packages/{pkg['id']}/sandbox-run", json={"plan_only": False}, headers=ADMIN).status_code == 409  # sandbox not configured
    assert client.get(f"/api/packages/{pkg['id']}", headers=ADMIN).json()["sandbox_report"]["mode"] == "plan"
    # D-05: merch lead sees cost, product merch does not
    lead = make_user(session, "lead@swimusa.com", ["merch_lead"])
    merch = make_user(session, "merch@swimusa.com", ["product_merch"])
    assert lead.can_view_group("cost") and lead.can("view_cost") and not merch.can_view_group("cost")

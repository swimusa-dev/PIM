"""IMP-01 to IMP-08, IMP-12: parse, map, validate, match, review, commit, rollback, worksheets."""

import io

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import select
from tests.conftest import FIXTURES

from pim.imports.batches import add_file, create_batch, save_template, set_mapping
from pim.imports.mapping import suggest_mappings
from pim.imports.matching import extract_name, match_batch
from pim.imports.parsers import detect_header, parse_upload, structure_signature
from pim.imports.survivorship import commit_batch, decide, decide_conflicts, review_batch, rollback_batch
from pim.imports.transforms import apply_one, combine_sources
from pim.imports.validation import error_file, save_value_mappings, validate_batch, value_mapping_report
from pim.imports.worksheet import build_worksheet
from pim.models import ImportConflict, ProductEntity, ProductValue
from pim.services.storage import Storage
from pim.services.values import set_value

SUPPLIER_MAPPING = [
    {"target": "style_number", "sources": ["Style #"], "transforms": []},
    {"target": "color_code", "sources": ["Colour"], "transforms": []},
    {"target": "upc", "sources": ["UPC"], "transforms": []},
    {"target": "product_name", "sources": ["Product Name"], "transforms": [{"op": "trim", "args": {}}]},
    {"target": "fabric_story_text", "sources": ["Fabric Story"], "transforms": [], "skip": True},
    {"target": "weight", "sources": ["Weight (lb)"], "transforms": []},
    {"target": "ecomm_launch_date", "sources": ["Launch Date"], "transforms": [{"op": "date_parse", "args": {}}]},
    {"target": "care_instructions", "sources": ["Care"], "transforms": []},
]


def _sc(session, style, color):
    return session.execute(select(ProductEntity).where(ProductEntity.level == "style_color", ProductEntity.style_number == style,
                                                       ProductEntity.color_code == color)).scalar_one()


def _pv(session, ent, code):
    return session.execute(select(ProductValue).where(ProductValue.entity_id == ent.id, ProductValue.attribute_code == code,
                                                      ProductValue.channel_code == "")).scalar_one_or_none()


def _supplier_batch(session, meta, admin, tmp_path):
    batch = create_batch(session, admin, "supplier", "style_color")
    add_file(session, meta, batch, "supplier_sample.csv", (FIXTURES / "supplier_sample.csv").read_bytes(), Storage(str(tmp_path)))
    return batch


# ---- parsing (IMP-01) -----------------------------------------------------------------------------------------
def test_imp01_setup_file_layout_two_header_rows_and_text_dates():
    files = parse_upload("setup_file_sample.csv", (FIXTURES / "setup_file_sample.csv").read_bytes())
    f = files[0]
    assert f.header_row == 2 and f.group_header_row == 1
    names = [c.name for c in f.columns]
    assert "Product Description" in names and f.columns[names.index("Fit Details")].group == "Design & Tech"
    assert f.rows[0]["Ecomm Launch Date"] == "10.30.26" and f.rows[0]["Fit Details"].startswith("• Removable")


def test_imp01_xlsx_merged_group_headers_and_header_not_on_row_1():
    wb = Workbook()
    ws = wb.active
    ws.append(["VA Resort 2027 setup file"])
    ws.append([])
    ws.append(["Style Details", None, "Copy"])
    ws.merge_cells("A3:B3")
    ws.append(["Style #", "Color", "Product Description"])
    ws.append(["200809", "FVP", "Meet the Skylar Top."])
    buf = io.BytesIO()
    wb.save(buf)
    f = parse_upload("setup.xlsx", buf.getvalue())[0]
    assert f.header_row == 4 and f.group_header_row == 3
    assert [c.group for c in f.columns] == ["Style Details", "Style Details", "Copy"]
    assert f.rows[0]["Style #"] == "200809"


def test_detect_header_and_signature():
    grid = [["Title only"], [], ["Style", "Color", "Size"], ["1", "A", "S"]]
    assert detect_header(grid) == (2, None)
    assert structure_signature(["Style", "Color"]) == structure_signature([" style", "COLOR "])


# ---- transforms (IMP-02) ---------------------------------------------------------------------------------------
def test_imp02_transforms():
    assert apply_one("  x ", "trim", {}) == "x"
    assert apply_one("a|b|c", "split", {"delimiter": "|", "index": 1}) == "b"
    assert apply_one("Front rise 3 1/4in", "regex_extract", {"pattern": r"rise\s+(.+)", "group": 1}) == "3 1/4in"
    assert apply_one("3 1/4", "unit_convert", {"from": "in", "to": "cm"}) == "8.255"
    assert apply_one("<p>Hello&nbsp;world</p>", "html_clean", {}) == "Hello world"
    assert apply_one("10.30.26", "date_parse", {}) == "2026-10-30"
    assert apply_one("FSTVL PINK", "case", {"mode": "title"}) == "Fstvl Pink"
    assert combine_sources({"a": "Skylar", "b": "Top"}, ["a", "b"], [{"op": "concat", "args": {"separator": " "}}]) == "Skylar Top"


def test_imp02_suggestions_hidden_below_threshold_and_never_auto_applied(meta):
    s = suggest_mappings(["Style #", "Colour", "Product Name", "Zebra Quotient"], meta, "style_color")
    by_col = {x["column"]: x for x in s}
    assert by_col["Style #"]["target"] == "style_number" and by_col["Style #"]["confidence"] >= 0.7
    assert by_col["Colour"]["target"] == "color_code"
    assert by_col["Product Name"]["target"] == "product_name"
    assert "Zebra Quotient" not in by_col


# ---- validate, value mappings, match, review, commit (IMP-03 to IMP-07) -------------------------------------
def test_import_pipeline_end_to_end(session, meta, admin, ingested, tmp_path):
    batch = _supplier_batch(session, meta, admin, tmp_path)
    assert batch.status == "uploaded" and len(batch.rows) == 4 and batch.files[0]["signature"]
    missing = set_mapping(session, meta, batch, SUPPLIER_MAPPING, ["upc", "style_color"], {"on_error": "skip_value"}, None, None)
    assert missing == []
    # IMP-03: incoming care values need mappings; the report suggests one
    rep = value_mapping_report(session, meta, batch)
    care = next(x for x in rep if x["attribute_code"] == "care_instructions")
    raws = {v["raw"]: v for v in care["values"]}
    assert raws["Hand Wash Cold Dry Flat"]["suggested"] == "hand_wash_cold_dry_flat"
    save_value_mappings(session, "supplier", [{"list_code": "care_instructions", "raw": "Hand Wash Cold Dry Flat", "option_code": "hand_wash_cold_dry_flat"},
                                              {"list_code": "care_instructions", "raw": "Hand wash cold. Line dry.", "option_code": "hand_wash_cold_line_dry"}], admin.email)
    # IMP-04 validation: placeholders removed, dates parsed, unmapped select values become errors (skip_value keeps the row)
    validate_batch(session, meta, batch)
    assert batch.stats["placeholders_removed"].get("care_instructions") == 1
    r1 = batch.rows[0]
    assert r1.mapped["ecomm_launch_date"] == "2026-10-30" and r1.mapped["weight"] == "0.16" and r1.mapped["care_instructions"] == "hand_wash_cold_dry_flat"
    assert r1.mapped["product_name"] == "Skylar Top"
    name, data = error_file(batch)
    assert name.endswith(".csv") and b"Errors" in data
    # IMP-05 matching: UPC first, then style+color; unmatched row stays unmatched under update_only
    stats = match_batch(session, meta, batch)
    assert stats["matched"] == 3 and stats["unmatched"] == 1
    assert batch.rows[0].match_key_used == "upc" and batch.rows[0].entity_id == _sc(session, "200809", "FVP").id
    assert batch.rows[3].action == "skip" and batch.rows[3].match_status == "unmatched"
    # IMP-06 survivorship: an approved user value on product_name is held as a conflict, empty values are filled
    sc = _sc(session, "200809", "FVP")
    st = session.get(ProductEntity, sc.parent_id)
    set_value(session, meta, admin, st, meta.attributes["product_name"], "Skylar Longline Top")
    session.flush()
    diff = review_batch(session, meta, batch)
    assert diff["conflicts"] >= 1 and any(a["attribute_code"] == "weight" for a in diff["adds"])
    assert batch.stats["skipped"] == 1
    conflict = session.execute(select(ImportConflict).where(ImportConflict.batch_id == batch.id, ImportConflict.attribute_code == "product_name")).scalars().first()
    assert conflict.current["value"] == "Skylar Longline Top" and conflict.incoming["value"] == "Skylar Top"
    # IMP-07 commit: conflicts held unless accepted; everything is one change set
    decide_conflicts(session, admin, [conflict.id], "reject")
    report = commit_batch(session, meta, admin, batch)
    assert batch.status == "committed" and report["values_approved"] >= 3 and report["conflicts_rejected"] >= 1
    assert _pv(session, st, "product_name").value_text == "Skylar Longline Top"
    assert _pv(session, sc, "weight").value_number == 0.16 and _pv(session, sc, "weight").source_type == "import"
    assert _pv(session, sc, "weight").source_ref.startswith("supplier:")
    # IMP-07 rollback restores prior values as a unit, unless touched later
    set_value(session, meta, admin, sc, meta.attributes["weight"], "0.5")
    session.flush()
    result = rollback_batch(session, meta, admin, batch)
    assert batch.status == "rolled_back" and result["restored"] >= 2
    assert any(m["attribute_code"] == "weight" for m in result["manual"])
    assert _pv(session, sc, "ecomm_launch_date") is None and _pv(session, sc, "weight").value_number == 0.5


def test_imp06_decision_matrix(session, meta, ingested):
    sc = _sc(session, "200809", "FVP")
    assert decide("never_from_import", None, 3) == "skip"
    assert decide("fill_empty_only", None, 3) == "write"
    pv = ProductValue(entity_level="style_color", entity_id=sc.id, attribute_code="x", value_text="v", source_type="user", approval_state="approved")
    assert decide("fill_empty_only", pv, 3) == "conflict"
    assert decide("trusted_source", pv, 0) == "write"  # ERP beats user
    assert decide("trusted_source", pv, 3) == "skip"  # supplier never replaces user-approved silently
    assert decide("most_recent", pv, 3) == "write"
    default = ProductValue(entity_level="style_color", entity_id=sc.id, attribute_code="x", value_text="v", source_type="default", confirmed_default=False)
    assert decide("fill_empty_only", default, 3) == "write"  # unconfirmed defaults count as empty


def test_imp05_create_policy_and_create_under_match(session, meta, admin, ingested, tmp_path):
    batch = _supplier_batch(session, meta, admin, tmp_path)
    set_mapping(session, meta, batch, SUPPLIER_MAPPING, ["upc", "style_color"], {"unmatched_policy": "create"}, None, None)
    validate_batch(session, meta, batch)
    stats = match_batch(session, meta, batch)
    assert stats["new"] == 1 and batch.rows[3].action == "create"
    review_batch(session, meta, batch)
    assert batch.stats["creates"] == 1
    report = commit_batch(session, meta, admin, batch)
    created = report["created_entities"]
    assert len(created) == 2  # style 999999 and style-color 999999-NEW
    new_sc = _sc(session, "999999", "NEW")
    assert new_sc.lifecycle_status == "new" and _pv(session, new_sc, "weight").value_number == 0.2
    rollback_batch(session, meta, admin, batch)
    assert session.execute(select(ProductEntity).where(ProductEntity.style_number == "999999")).first() is None


def test_imp05_keyless_name_extraction_and_candidates(session, meta, admin, ingested, tmp_path):
    assert extract_name("Meet the Skylar Top, a longline bralette in our signature EcoRib fabric.") == "Skylar Top"
    batch = create_batch(session, admin, "other", "style_color")
    add_file(session, meta, batch, "setup_file_sample.csv", (FIXTURES / "setup_file_sample.csv").read_bytes(), Storage(str(tmp_path)))
    set_mapping(session, meta, batch, [{"target": "product_description", "sources": ["Product Description"], "transforms": []}], [],
                {"unmatched_policy": "update_only", "keyless_name_column": "Product Description"}, None, None)
    validate_batch(session, meta, batch)
    assert any("assisted matching" in w["message"] for w in batch.rows[0].warnings)
    match_batch(session, meta, batch)
    row = batch.rows[0]
    assert row.match_status == "unmatched" and row.mapped["_extracted_name"] == "Skylar Top"
    assert row.candidates and "SKYLAR" in row.candidates[0]["label"].upper()
    from pim.imports.matching import confirm_row

    confirm_row(session, row, row.candidates[0]["entity_id"], "update", admin.email)
    assert row.match_status == "matched" and row.match_key_used == "assisted" and row.decided_by == admin.email


def test_imp02_template_saved_and_auto_applied(session, meta, admin, ingested, tmp_path):
    batch = _supplier_batch(session, meta, admin, tmp_path)
    set_mapping(session, meta, batch, SUPPLIER_MAPPING, ["upc"], {}, None, None)
    tpl = save_template(session, batch, "Vendor A standard", admin)
    session.flush()
    batch2 = _supplier_batch(session, meta, admin, tmp_path)
    assert batch2.status == "mapped" and batch2.template_id == tpl.id and batch2.match_keys == ["upc"]


def test_imp04_skip_row_policy_and_erp_protection(session, meta, admin, ingested, tmp_path):
    batch = _supplier_batch(session, meta, admin, tmp_path)
    mapping = [*SUPPLIER_MAPPING, {"target": "erp_style_content", "sources": ["Fabric Story"], "transforms": []},
               {"target": "weight", "sources": ["Weight (lb)"], "transforms": [], "required": True}]
    set_mapping(session, meta, batch, mapping, ["upc"], {"on_error": "skip_row"}, None, None)
    validate_batch(session, meta, batch)
    assert any("ERP-owned" in e["message"] for e in batch.rows[0].errors)
    assert batch.rows[0].action == "skip" and batch.stats["errors"] >= 1


# ---- keyed worksheets (IMP-12) ----------------------------------------------------------------------------------
def test_imp12_worksheet_round_trip(session, meta, admin, ingested, tmp_path):
    ents = session.execute(select(ProductEntity).where(ProductEntity.level == "style_color", ProductEntity.style_number == "200809")).scalars().all()
    ws, data = build_worksheet(session, meta, admin, ents, "style_color", ["style_details", "seo"], True, {"entity_ids": [e.id for e in ents]})
    ws.path = Storage(str(tmp_path)).write(f"ws/{ws.id}.xlsx", data)
    session.flush()
    wb = load_workbook(io.BytesIO(data))
    sheet = wb["Product Information"]
    assert sheet.freeze_panes == "G3" and sheet.protection.sheet and sheet.column_dimensions["A"].hidden
    assert sheet["A1"].value == "Key" and sheet["B2"].value == "_style_number"
    assert "_pim" in wb.sheetnames and wb["_pim"].sheet_state == "hidden" and "Instructions" in wb.sheetnames
    headers = [c.value for c in sheet[2]]
    col = headers.index("Customer-facing color name") + 1
    sheet.cell(row=3, column=col, value="Festival Pink")
    sheet.cell(row=4, column=col, value="Eco Black")
    sheet.cell(row=4, column=2, value="TAMPERED")  # edit a locked key column
    buf = io.BytesIO()
    wb.save(buf)
    batch = create_batch(session, admin, "worksheet", "style_color")
    add_file(session, meta, batch, ws.file_name, buf.getvalue(), Storage(str(tmp_path)))
    assert batch.worksheet_id == ws.id and batch.status == "mapped" and batch.match_keys == ["_pim_id"]
    validate_batch(session, meta, batch)
    rows = sorted(batch.rows, key=lambda r: r.row_number)
    assert rows[0].match_status == "matched" and rows[0].mapped["customer_color_name"] == "Festival Pink"
    assert rows[1].action == "skip" and any("Locked key columns" in e["message"] for e in rows[1].errors)
    match_batch(session, meta, batch)
    review_batch(session, meta, batch)
    commit_batch(session, meta, admin, batch)
    assert _pv(session, rows[0].entity_id and session.get(ProductEntity, rows[0].entity_id), "customer_color_name").value_text == "Festival Pink"


def test_imp08_shopify_rows(session, meta, admin, ingested):
    from pim.adapters.shopify_admin import parse_product_csv
    from pim.imports.batches import rows_from_shopify

    batch = create_batch(session, admin, "shopify", "sku")
    snap = parse_product_csv((FIXTURES / "shopify_products_export.csv").read_bytes())
    n = rows_from_shopify(session, batch, snap)
    session.flush()
    assert n == 7 and batch.rows[0].raw["Handle"] == "skylar-top-festival-pink" and "metafield:custom.care_instructions" in batch.rows[0].raw
    set_mapping(session, meta, batch, [{"target": "sku", "sources": ["Variant SKU"], "transforms": []},
                                       {"target": "shopify_tags", "sources": ["Tags"], "transforms": []}], ["sku"], {}, None, None)
    validate_batch(session, meta, batch)
    stats = match_batch(session, meta, batch)
    assert stats["matched"] == 7


@pytest.mark.parametrize("policy,expected", [("fill_empty_only", "conflict"), ("most_recent", "write")])
def test_imp06_group_policy_override(meta, policy, expected):
    from pim.imports.survivorship import policy_for

    attr = meta.attributes["seo_title"]
    attr.survivorship = policy
    try:
        assert policy_for(meta, attr) == policy
    finally:
        attr.survivorship = None
    assert policy_for(meta, meta.attributes["erp_style_content"]) == "never_from_import"

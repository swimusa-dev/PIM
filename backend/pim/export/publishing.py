"""EXP-05 publish history and feedback loop, EXP-06 delta and scheduled exports.

A PublishRecord is created for every dry run, package, publish and scheduled run with its channel version, the
selection, validation counts, the files written to the publish volume and one PublishItem per product. Retailer
feedback (an error report file or an adapter poll) is mapped back to attributes through the channel mapping and
tagged `internal` (our data) or `external` (retailer side), item statuses move, and rejections notify the users
holding the channel's owning role. Health score = accepted / (accepted + rejected) over the last 30 days.

Sending to a retailer API is behind PIM_FEATURE_CHANNEL_SEND; EXP-07 direct Shopify writes stay behind
PIM_FEATURE_SHOPIFY_DIRECT_WRITE, so the Shopify publish path produces an EXP-13 package (Phase 1 decision).
"""

from __future__ import annotations

import csv
import io
import json
import logging
import re
from collections import Counter
from datetime import UTC, datetime, timedelta
from typing import Any

from openpyxl import load_workbook
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from pim.adapters.base import NotConfigured
from pim.auth.identity import CurrentUser
from pim.config import get_settings
from pim.export.payloads import ProductPayload, build_payloads, payload_files
from pim.export.schema_loaders import TargetSchema
from pim.export.shopify_package import resolve_selection
from pim.models import (
    AssetLink,
    Channel,
    EligibilityOverride,
    ProductEntity,
    ProductValue,
    PublishItem,
    PublishRecord,
    ReadinessSnapshot,
    UserRole,
)
from pim.models._common import utcnow
from pim.services.collab import notify
from pim.services.metadata import Metadata
from pim.services.storage import Storage
from pim.services.telemetry import record_adapter_error

log = logging.getLogger(__name__)

TERMINAL = ("completed", "failed", "skipped", "submitted", "accepted", "partial", "rejected")
EXTERNAL_HINTS = re.compile(r"\b(internal server|timeout|rate limit|throttl|unavailable|retailer|marketplace|category not|"
                            r"setup by match|item not found|gtin.*(already|exists|in use)|duplicate|restricted brand|"
                            r"brand not authorized|prohibited|not allowed to sell|invalid request|service)\b", re.IGNORECASE)


def storage() -> Storage:
    s = get_settings()
    return Storage(getattr(s, "publish_volume", None) or s.package_volume)


# ---------------------------------------------------------------------------------------------------
# Selection and delta (EXP-06)
# ---------------------------------------------------------------------------------------------------
def last_success_at(session: Session, channel_code: str) -> datetime | None:
    return session.execute(select(func.max(PublishRecord.finished_at)).where(
        PublishRecord.channel_code == channel_code, PublishRecord.kind.in_(["publish", "scheduled", "package"]),
        PublishRecord.status.in_(["completed", "submitted", "accepted", "partial"]))).scalar_one_or_none()


def eligible_entities(session: Session, meta: Metadata, channel: Channel, entities: list[ProductEntity]) -> tuple[list[ProductEntity], list[dict]]:
    """Drop ineligible products (EXP-12) unless an unexpired override exists. Returns (kept, skipped)."""
    from pim.services.readiness import check_eligibility

    now = utcnow()
    overrides = {o.entity_id: o for o in session.execute(select(EligibilityOverride).where(
        EligibilityOverride.channel_code == channel.code)).scalars() if o.expires_at is None or o.expires_at > now}
    styles = {s.id: s for s in session.execute(select(ProductEntity).where(
        ProductEntity.id.in_({e.parent_id for e in entities if e.parent_id}))).scalars()} if entities else {}
    kept, skipped = [], []
    for e in entities:
        reasons = check_eligibility(channel, e, styles.get(e.parent_id or ""))
        if reasons and e.id not in overrides:
            skipped.append({"entity_id": e.id, "label": e.label, "reasons": reasons})
        else:
            kept.append(e)
    return kept, skipped


def delta_candidates(session: Session, meta: Metadata, channel: Channel, base: list[ProductEntity] | None = None) -> list[dict]:
    """Products whose values, images or readiness changed since the last successful publish to the channel."""
    since = last_success_at(session, channel.code)
    if base is None:
        base = session.execute(select(ProductEntity).where(ProductEntity.level == "style_color")).scalars().all()
    base, _ = eligible_entities(session, meta, channel, base)
    ids = [e.id for e in base]
    if not ids:
        return []
    by_id = {e.id: e for e in base}
    parents: dict[str, list[str]] = {}
    for e in base:
        if e.parent_id:
            parents.setdefault(e.parent_id, []).append(e.id)  # a style change touches every colorway under it
    reasons: dict[str, list[str]] = {}
    if since is None:
        return [{"entity_id": e.id, "label": e.label, "reasons": ["Never published to this channel"]} for e in base]
    for pv in session.execute(select(ProductValue.entity_id, ProductValue.attribute_code).where(
            ProductValue.entity_id.in_(ids + list(parents)), ProductValue.set_at >= since, ProductValue.approval_state == "approved")):
        targets = [pv.entity_id] if pv.entity_id in by_id else parents.get(pv.entity_id, [])
        for target in targets:
            reasons.setdefault(target, []).append(f"{pv.attribute_code} changed")
    for lk in session.execute(select(AssetLink).where(AssetLink.entity_id.in_(ids), AssetLink.decided_at >= since)).scalars():
        reasons.setdefault(lk.entity_id, []).append(f"image {lk.shot_type or lk.asset_id} {lk.status}")
    for snap in session.execute(select(ReadinessSnapshot).where(ReadinessSnapshot.entity_id.in_(ids),
                                                                ReadinessSnapshot.channel_code == channel.code,
                                                                ReadinessSnapshot.computed_at >= since)).scalars():
        reasons.setdefault(snap.entity_id, []).append("readiness recomputed" if snap.ready else "readiness changed (not ready)")
    sent_before = {r for r in session.execute(select(PublishItem.entity_id).join(PublishRecord).where(
        PublishRecord.channel_code == channel.code, PublishItem.status.in_(["submitted", "accepted", "in_review"]))).scalars()}
    for e in base:
        if e.id not in sent_before:
            reasons.setdefault(e.id, []).append("Never sent to this channel")
    return [{"entity_id": eid, "label": by_id[eid].label, "reasons": sorted(set(r))} for eid, r in reasons.items()]


# ---------------------------------------------------------------------------------------------------
# Runs (EXP-04 dry run, EXP-05 publish)
# ---------------------------------------------------------------------------------------------------
def start_run(session: Session, channel: Channel, kind: str, mode: str, user_email: str, selection: dict) -> PublishRecord:
    rec = PublishRecord(channel_code=channel.code, channel_version=channel.version, kind=kind, mode=mode, status="running",
                        created_by=user_email, selection=selection or {})
    session.add(rec)
    session.flush()
    return rec


def running_run(session: Session, channel_code: str) -> PublishRecord | None:
    return session.execute(select(PublishRecord).where(PublishRecord.channel_code == channel_code,
                                                       PublishRecord.status == "running").order_by(PublishRecord.created_at.desc())).scalars().first()


def execute_run(session: Session, meta: Metadata, channel: Channel, rec: PublishRecord, *, preview: int = 5,
                send: bool = False, user: CurrentUser | None = None) -> PublishRecord:
    """Resolve, filter, build, validate, write files, store items; for publishes also send (when allowed)."""
    st = storage()
    try:
        entities = resolve_selection(session, rec.selection or {})
        if rec.mode == "delta":
            # The selection is the ceiling: an empty selection is an empty run, never the whole catalogue.
            delta = {d["entity_id"] for d in delta_candidates(session, meta, channel, entities)}
            entities = [e for e in entities if e.id in delta]
        entities, skipped = eligible_entities(session, meta, channel, entities)
        schema = TargetSchema.from_dict(channel.target_schema)
        rends = _renditions_for(session, channel, entities)
        payloads = build_payloads(session, meta, channel, entities, schema, rends)
        valid = [p for p in payloads if p.valid]
        rec.product_count, rec.valid_count, rec.invalid_count = len(payloads), len(valid), len(payloads) - len(valid)
        files = payload_files(channel, payloads, schema, template_bytes_for(channel))
        rec.payload_path = st.path("publishes", rec.id)
        rec.files = []
        for name, content, kind in files:
            full = st.write(f"publishes/{rec.id}/{name}", content)
            rec.files.append({"name": name, "size": len(content), "kind": kind, "path": full})
        issue_counter = Counter(i.kind for p in payloads for i in p.issues if i.severity == "error")
        rec.summary = {"preview": [{"entity_id": p.entity_id, "label": p.label, "external_id": p.external_id, "data": p.data,
                                    "images": p.images} for p in payloads[: max(0, preview)]],
                       "skipped_ineligible": skipped[:200], "skipped_ineligible_count": len(skipped),
                       "issues_by_kind": dict(issue_counter), "schema_source": schema.source, "schema_fields": len(schema.fields),
                       "mode": rec.mode}
        for p in payloads:
            session.add(PublishItem(publish_id=rec.id, entity_id=p.entity_id, external_id=p.external_id,
                                    status="valid" if p.valid else "invalid", payload=p.data,
                                    issues=[i.as_dict() for i in p.issues], errors=[]))
        session.flush()
        if rec.kind in ("publish", "scheduled") and valid:
            _send(session, meta, channel, rec, valid, send, user)
        else:
            rec.status = "completed"
    except Exception as exc:  # noqa: BLE001
        log.exception("publish run %s failed", rec.id)
        rec.status, rec.error = "failed", str(exc)[:2000]
        record_adapter_error(session, channel.adapter or channel.code, exc, {"publish_id": rec.id, "channel": channel.code, "kind": rec.kind})
    rec.finished_at = utcnow()
    return rec


def _renditions_for(session: Session, channel: Channel, entities: list[ProductEntity]) -> dict[str, list[dict]]:
    from pim.models import ImageRendition

    if not entities:
        return {}
    asset_ids = {lk.asset_id for lk in session.execute(select(AssetLink).where(
        AssetLink.entity_id.in_([e.id for e in entities]), AssetLink.status == "confirmed")).scalars()}
    out: dict[str, list[dict]] = {}
    if asset_ids:
        for r in session.execute(select(ImageRendition).where(ImageRendition.asset_id.in_(asset_ids),
                                                              ImageRendition.channel_code == channel.code)).scalars():
            out.setdefault(r.asset_id, []).append({"channel_code": r.channel_code, "url": r.url, "width": r.width, "height": r.height})
    return out


SEND_DISABLED = "send disabled (PIM_FEATURE_CHANNEL_SEND=false)"


def send_enabled() -> bool:
    """EXP-05/EXP-09: one switch for every outbound retailer send (API feeds, SFTP uploads, emailed files)."""
    return bool(getattr(get_settings(), "feature_channel_send", False))


def _send(session: Session, meta: Metadata, channel: Channel, rec: PublishRecord, valid: list[ProductPayload],
          send: bool, user: CurrentUser | None) -> None:
    transport = (channel.config or {}).get("transport") or ("package" if channel.adapter == "shopify_admin" else "download")
    items = {i.entity_id: i for i in session.execute(select(PublishItem).where(PublishItem.publish_id == rec.id)).scalars()}
    if transport == "package" or channel.adapter == "shopify_admin":
        # EXP-07 stays off: the Shopify publish is an EXP-13 package built from the same selection.
        from pim.models import UpdatePackage

        pkg = UpdatePackage(channel_code=channel.code, created_by=rec.created_by, status="pending",
                            selection={"entity_ids": [p.entity_id for p in valid]}, snapshot_mode=(channel.config or {}).get("snapshot_mode", "api"))
        session.add(pkg)
        session.flush()
        rec.package_id = pkg.id
        rec.summary = {**(rec.summary or {}), "package_id": pkg.id}  # reassign: plain JSON columns do not track in-place edits
        rec.status = "completed"
        for p in valid:
            items[p.entity_id].status = "submitted"
        return
    if transport in ("sftp", "email"):
        # EXP-09: file delivery through the channel transport; a failed delivery keeps the publish and its files.
        # PIM_FEATURE_CHANNEL_SEND gates every outbound retailer send, file deliveries included.
        rec.status = "completed"
        if not send:
            result = {"transport": transport, "status": "skipped", "detail": "dry run", "at": utcnow().isoformat(), "files": []}
            rec.summary = {**(rec.summary or {}), "delivery": result}
        elif not getattr(get_settings(), "feature_channel_send", False):
            result = {"transport": transport, "status": "skipped", "detail": SEND_DISABLED, "at": utcnow().isoformat(), "files": []}
            rec.summary = {**(rec.summary or {}), "delivery": result}
        else:
            result = deliver_publish(session, channel, rec)
        for p in valid:
            items[p.entity_id].status = "submitted" if result.get("status") == "delivered" else items[p.entity_id].status
        return
    if transport == "download" or not send or not getattr(get_settings(), "feature_channel_send", False):
        rec.status = "completed"
        rec.summary = {**(rec.summary or {}), "delivery": {"transport": transport, "status": "skipped", "at": utcnow().isoformat(), "files": [],
                                                           "detail": "download transport" if transport == "download" else SEND_DISABLED}}
        return
    adapter = adapter_for(channel)
    try:
        if channel.adapter == "walmart_mp":
            res = _send_walmart(channel, adapter, valid)
        else:
            res = adapter.publish([p.data for p in valid], dry_run=False)
    except NotConfigured as exc:
        rec.status, rec.error = "failed", str(exc)
        record_adapter_error(session, channel.adapter or channel.code, exc, {"publish_id": rec.id, "channel": channel.code})
        return
    rec.external_batch_id = res.batch_id
    rec.status = "submitted"
    now = utcnow()
    rec.summary = {**(rec.summary or {}), "delivery": {"transport": "api", "status": "delivered", "at": now.isoformat(), "files": [],
                                                       "detail": f"batch {res.batch_id}: {res.accepted} item(s) submitted", "batch_id": res.batch_id,
                                                       "accepted": res.accepted, "rejected": res.rejected, **res.detail}}
    if channel.adapter == "walmart_mp":
        from pim.adapters.walmart_mp import next_poll_at

        rec.summary = {**rec.summary, "walmart": {"feed_ids": res.detail.get("feed_ids", []), "status": "submitted", "polls": 0,
                                                  "next_poll_at": next_poll_at(0, now).isoformat(), "spec_version": res.detail.get("spec_version")}}
    for p in valid:
        items[p.entity_id].status = "submitted"


def _send_walmart(channel: Channel, adapter, valid: list[ProductPayload]):  # noqa: ANN001, ANN202
    """EXP-08: nest each payload into Orderable and Visible, group variants per style [C12], validate, submit."""
    from pim.export.payloads import nest
    from pim.export.schema_loaders import TargetSchema

    schema = TargetSchema.from_dict(channel.target_schema)
    w = (channel.config or {}).get("walmart") or {}
    items = []
    for p in valid:
        nested = nest(p.data, schema)
        # The variant group is the style (every color of a style is one Walmart variant group); the style number is
        # carried on the payload because external ids are "{style}-{color}" and a style number may itself hold a hyphen.
        items.append({"orderable": nested.get("Orderable", {}), "visible": nested.get("Visible", {}),
                      "variant_group": p.style_number or (p.external_id.split("-")[0] if p.external_id else None), "_entity_id": p.entity_id})
    adapter.group_variants(items, w.get("variant_attribute_names") or [])
    check = adapter.validate(items)
    if not check.ok:
        raise NotConfigured("Walmart variant group validation failed: " + "; ".join(e["message"] for e in check.errors[:5]))
    return adapter.publish(items, dry_run=False)


def adapter_for(channel: Channel):  # noqa: ANN201
    cfg = channel.config or {}
    if channel.adapter == "walmart_mp":
        from pim.adapters.walmart_mp import WalmartMarketplaceAdapter

        w = cfg.get("walmart") or {}
        return WalmartMarketplaceAdapter(spec_version=str(w.get("spec_version") or cfg.get("spec_version") or "5.0"),
                                         max_items_per_feed=int(w.get("max_items_per_feed") or cfg.get("max_items_per_feed") or 5000))
    if channel.adapter == "rithum":
        from pim.adapters.rithum import rithum_for_channel

        return rithum_for_channel(channel)
    if channel.adapter == "mirakl_seller":
        from pim.adapters.mirakl_seller import MiraklSellerAdapter

        return MiraklSellerAdapter()
    if channel.adapter == "target_plus":
        from pim.adapters.target_plus import TargetPlusAdapter

        return TargetPlusAdapter()
    raise NotConfigured(f"No API adapter for {channel.adapter}")


def template_bytes_for(channel: Channel) -> bytes | None:
    path = (channel.config or {}).get("feed_template_path")
    if not path:
        return None
    try:
        return storage().read(path)
    except Exception as exc:  # noqa: BLE001
        log.warning("feed template for %s unreadable at %s: %s", channel.code, path, exc)
        return None


def deliver_publish(session: Session, channel: Channel, rec: PublishRecord) -> dict:
    """EXP-09: send the publish's payload files through the channel transport (SFTP or email) and record the result.
    Refuses (recorded as skipped) while PIM_FEATURE_CHANNEL_SEND is off."""
    from pim.services.delivery import deliver

    if not send_enabled():
        result = {"transport": (channel.config or {}).get("transport") or "download", "status": "skipped", "detail": SEND_DISABLED,
                  "at": utcnow().isoformat(), "files": []}
        rec.summary = {**(rec.summary or {}), "delivery": result}
        return result
    st = storage()
    files = [(f["name"], st.read(f["path"])) for f in (rec.files or []) if f.get("kind") == "payload" and f.get("path")]
    result = deliver(channel, files)
    rec.summary = {**(rec.summary or {}), "delivery": result}
    if result.get("status") == "failed":
        record_adapter_error(session, f"{channel.code}:{result.get('transport')}", result.get("detail", ""), {"publish_id": rec.id, "channel": channel.code})
    return result


# ---------------------------------------------------------------------------------------------------
# Feedback (EXP-05)
# ---------------------------------------------------------------------------------------------------
ID_KEYS = ("external_id", "sku", "product id", "productid", "item id", "itemid", "product_id", "id", "style", "style_number", "handle", "gtin", "upc")
MSG_KEYS = ("message", "error", "errors", "description", "reason", "ingestion errors", "error message", "detail", "issue")
FIELD_KEYS = ("field", "attribute", "attribute code", "attribute_code", "column", "property", "path")
STATUS_KEYS = ("status", "state", "result", "ingestion status")
ACCEPTED_WORDS = ("accepted", "success", "successful", "published", "ok", "live", "active", "processed")
REJECTED_WORDS = ("rejected", "error", "errors", "failed", "fail", "unpublished")
REVIEW_WORDS = ("in_review", "in review", "review", "pending", "processing", "in progress", "in_progress")


def parse_feedback_file(name: str, content: bytes) -> list[dict]:
    """Rows of {external_id, message, field?, status?} from CSV, JSON or XLSX retailer error reports."""
    rows: list[dict] = []
    lower = name.lower()
    if lower.endswith(".json") or content[:1] in (b"{", b"["):
        doc = json.loads(content.decode("utf-8-sig"))
        if isinstance(doc, dict):
            doc = doc.get("errors") or doc.get("items") or doc.get("itemDetails", {}).get("itemIngestionStatus") or doc.get("data") or [doc]
        for it in doc:
            if isinstance(it, dict):
                rows.append(_norm_row(it))
        return rows
    if lower.endswith(".xlsx"):
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        ws = wb.worksheets[0]
        it = ws.iter_rows(values_only=True)
        header = [str(h or "").strip() for h in next(it, [])]
        for r in it:
            if any(c not in (None, "") for c in r):
                rows.append(_norm_row(dict(zip(header, r, strict=False))))
        wb.close()
        return rows
    text = content.decode("utf-8-sig", errors="replace")
    for r in csv.DictReader(io.StringIO(text)):
        rows.append(_norm_row(r))
    return rows


def _norm_row(r: dict) -> dict:
    low = {str(k).strip().lower(): v for k, v in r.items()}

    def pick(keys: tuple[str, ...]) -> Any:
        for k in keys:
            if k in low and low[k] not in (None, ""):
                return low[k]
        return None

    msg = pick(MSG_KEYS)
    if isinstance(msg, list):
        msg = "; ".join(str(m.get("description", m) if isinstance(m, dict) else m) for m in msg)
    status = pick(STATUS_KEYS)
    return {"external_id": str(pick(ID_KEYS) or "").strip(), "message": str(msg or "").strip(),
            "field": (str(pick(FIELD_KEYS)).strip() if pick(FIELD_KEYS) else None),
            "status": str(status).strip().lower() if status else None}


def classify_error(message: str, field_name: str | None, channel: Channel) -> tuple[str, str | None]:
    """Tag a retailer message internal (our data, mappable to an attribute) or external (retailer side)."""
    mapping: dict[str, str] = channel.attribute_mapping or {}
    rev = {v: k.split(".", 1)[0] for k, v in mapping.items() if v}
    attr = None
    if field_name:
        attr = rev.get(field_name) or rev.get(field_name.split("/")[-1])
    if attr is None:
        for target, code in rev.items():
            if target and re.search(rf"\b{re.escape(target)}\b", message or "", re.IGNORECASE):
                attr = code
                break
    if attr:
        return "internal", attr
    if EXTERNAL_HINTS.search(message or ""):
        return "external", None
    if field_name:
        return "internal", None
    return "external", None


def apply_feedback(session: Session, meta: Metadata, channel: Channel, rec: PublishRecord, rows: list[dict],
                   source: str = "file") -> dict[str, int]:
    items = session.execute(select(PublishItem).where(PublishItem.publish_id == rec.id)).scalars().all()
    by_ext = {(i.external_id or "").lower(): i for i in items}
    ents = {e.id: e for e in session.execute(select(ProductEntity).where(ProductEntity.id.in_([i.entity_id for i in items]))).scalars()} if items else {}
    by_sku: dict[str, PublishItem] = {}
    for i in items:
        e = ents.get(i.entity_id)
        if e:
            by_sku[e.style_number.lower()] = i
            if e.color_code:
                by_sku[f"{e.style_number}-{e.color_code}".lower()] = i
                by_sku[f"{e.style_number}{e.color_code}".lower()] = i
    counts = Counter()
    touched: set[str] = set()
    for r in rows:
        key = (r.get("external_id") or "").lower()
        item = by_ext.get(key) or by_sku.get(key) or by_sku.get(key.replace("_", "-"))
        if item is None:
            counts["unmatched"] += 1
            continue
        msg = r.get("message") or ""
        status = (r.get("status") or "").lower()
        # The retailer's status wins; a message only decides when no status says otherwise (an accepted row can
        # carry an informational message and must stay accepted).
        if status in ACCEPTED_WORDS:
            item.status = "accepted"
        elif status in REVIEW_WORDS:
            item.status = "in_review"
        elif status in REJECTED_WORDS or msg:
            item.status = "rejected"
        if msg and item.status != "accepted":
            kind, attr = classify_error(msg, r.get("field"), channel)
            item.errors = [*(item.errors or []), {"message": msg[:500], "field": r.get("field"), "attribute_code": attr, "kind": kind,
                                                  "source": source, "at": utcnow().isoformat()}]
        touched.add(item.id)
    for item in items:
        if item.id in touched:
            counts[item.status] += 1
    if touched:
        statuses = Counter(i.status for i in items)
        if statuses.get("rejected") and (statuses.get("accepted") or statuses.get("submitted")):
            rec.status = "partial"
        elif statuses.get("rejected") and not statuses.get("accepted"):
            rec.status = "rejected"
        elif statuses.get("accepted") and not statuses.get("rejected") and not statuses.get("in_review"):
            rec.status = "accepted"
        rec.summary = {**(rec.summary or {}), "feedback": {"at": utcnow().isoformat(), "source": source, **{k: v for k, v in counts.items()}}}
        rejected = [i for i in items if i.id in touched and i.status == "rejected"]
        if rejected:
            _notify_owners(session, channel, rec, rejected, ents)
    return {"updated": len(touched), "accepted": counts.get("accepted", 0), "rejected": counts.get("rejected", 0),
            "in_review": counts.get("in_review", 0), "unmatched": counts.get("unmatched", 0)}


def poll_feedback(session: Session, meta: Metadata, channel: Channel, rec: PublishRecord, force: bool = True) -> dict[str, int]:
    if not rec.external_batch_id:
        raise NotConfigured("This publish has no retailer batch id to poll; upload the retailer's error report instead")
    adapter = adapter_for(channel)
    if channel.adapter == "walmart_mp":
        return _poll_walmart(session, meta, channel, adapter, rec, force)
    status = adapter.fetch_status(rec.external_batch_id)
    errors = adapter.fetch_errors(rec.external_batch_id)
    rows = [_norm_row(e) for e in errors]
    for it in status.get("items", []) if isinstance(status, dict) else []:
        rows.append(_norm_row(it))
    return apply_feedback(session, meta, channel, rec, rows, source="adapter")


def feedback_pending(rec: PublishRecord) -> bool:
    """True while the retailer can still report on this publish: submitted, or a Walmart feed whose items are
    still processing (the record itself may already read partial or rejected after the first poll)."""
    if not rec.external_batch_id:
        return False
    if rec.status == "submitted":
        return True
    return ((rec.summary or {}).get("walmart") or {}).get("status") == "in_progress"


def walmart_poll_due(rec: PublishRecord, now: datetime | None = None) -> bool:
    w = (rec.summary or {}).get("walmart") or {}
    nxt = w.get("next_poll_at")
    if w.get("status") == "complete":
        return False
    return not nxt or datetime.fromisoformat(nxt) <= (now or utcnow())


def _poll_walmart(session: Session, meta: Metadata, channel: Channel, adapter, rec: PublishRecord, force: bool) -> dict[str, int]:  # noqa: ANN001
    """EXP-08 [C14][C15]: poll every feed of the publish, store item statuses and ingestion errors (a PROCESSED feed
    can still hold INPROGRESS items), schedule the next poll on Walmart's cadence, and pull Listing Quality scores."""
    from pim.adapters.walmart_mp import next_poll_at

    w = dict((rec.summary or {}).get("walmart") or {"feed_ids": [rec.external_batch_id], "polls": 0})
    now = utcnow()
    if not force and not walmart_poll_due(rec, now):
        return {"updated": 0, "accepted": 0, "rejected": 0, "in_review": 0, "unmatched": 0, "skipped": "not due"}
    rows: list[dict] = []
    all_complete = True
    feed_status: dict[str, str] = {}
    feed_errors: dict[str, str] = {}
    publish_items = session.execute(select(PublishItem).where(PublishItem.publish_id == rec.id)).scalars().all()
    for fid in w.get("feed_ids") or [rec.external_batch_id]:
        st = adapter.fetch_status(fid)
        feed_status[fid] = st.get("status", "")
        all_complete = all_complete and bool(st.get("complete"))
        for it in st.get("items", []):
            rows.append({"external_id": it.get("external_id"), "status": it.get("status"), "message": it.get("message") or "",
                         "field": it.get("field")})
        if st.get("error") and not st.get("items"):
            # Walmart rejected the whole feed (feedStatus ERROR, no item details): every item still waiting on it is
            # rejected with the feed-level reason, so the publish settles instead of polling forever.
            feed_errors[fid] = str(st["error"])
            for i in publish_items:
                if i.external_id and i.status in ("submitted", "in_review", "pending", "valid"):
                    rows.append({"external_id": i.external_id, "status": "rejected", "message": f"feed {fid} rejected: {st['error']}", "field": None})
    result = apply_feedback(session, meta, channel, rec, rows, source="walmart")
    w["polls"] = int(w.get("polls", 0)) + 1
    w["feed_status"] = feed_status
    if feed_errors:
        w["feed_errors"] = feed_errors
    w["last_poll_at"] = now.isoformat()
    w["status"] = "complete" if all_complete else "in_progress"
    w["next_poll_at"] = None if all_complete else next_poll_at(w["polls"], now).isoformat()
    skus = [i.external_id for i in publish_items if i.external_id]
    scores = adapter.listing_quality(skus) if skus else {}
    rec.summary = {**(rec.summary or {}), "walmart": w, "listing_quality": {**((rec.summary or {}).get("listing_quality") or {}), **scores}}
    return result


def _notify_owners(session: Session, channel: Channel, rec: PublishRecord, rejected: list[PublishItem], ents: dict) -> None:
    role = (channel.config or {}).get("owner_role") or "ecommerce"
    emails = {u.email for u in session.execute(select(UserRole).where(UserRole.role_code == role)).scalars()}
    emails.add(rec.created_by)
    labels = ", ".join(ents[i.entity_id].label for i in rejected[:5] if i.entity_id in ents)
    for email in emails:
        if email and email != "system":
            notify(session, email, "publish_rejected", f"{channel.label}: {len(rejected)} product(s) rejected",
                   f"Rejected: {labels}{' and more' if len(rejected) > 5 else ''}", f"/exports/{rec.id}")


# ---------------------------------------------------------------------------------------------------
# Health (EXP-05)
# ---------------------------------------------------------------------------------------------------
def channel_health(session: Session, channel: Channel, days: int = 30) -> dict:
    since = utcnow() - timedelta(days=days)
    recs = session.execute(select(PublishRecord).where(PublishRecord.channel_code == channel.code,
                                                       PublishRecord.kind.in_(["publish", "scheduled", "package"]),
                                                       PublishRecord.created_at >= since)).scalars().all()
    ids = [r.id for r in recs]
    items = session.execute(select(PublishItem).where(PublishItem.publish_id.in_(ids))).scalars().all() if ids else []
    st = Counter(i.status for i in items)
    accepted, rejected = st.get("accepted", 0), st.get("rejected", 0)
    score = round(accepted / (accepted + rejected), 3) if accepted + rejected else None
    errs = Counter()
    kinds: dict[str, str] = {}
    for i in items:
        for e in i.errors or []:
            errs[e["message"]] += 1
            kinds[e["message"]] = e.get("kind", "external")
    last = max((r.created_at for r in recs), default=None)
    out: dict[str, Any] = {}
    if channel.adapter == "walmart_mp":  # EXP-08 only: other channels carry neither key, so the UI shows no Walmart cards
        lq = [v for r in recs for v in ((r.summary or {}).get("listing_quality") or {}).values() if isinstance(v, (int, float))]
        out["listing_quality_avg"] = round(sum(lq) / len(lq), 1) if lq else None
        out["feeds_in_progress"] = sum(1 for r in recs if ((r.summary or {}).get("walmart") or {}).get("status") == "in_progress")
    return {"channel_code": channel.code, "publishes": len(recs), "products_sent": sum(r.product_count for r in recs), **out,
            "accepted": accepted, "rejected": rejected, "in_review": st.get("in_review", 0), "health_score": score,
            "last_publish_at": last.isoformat() if last else None,
            "last_success_at": (lambda d: d.isoformat() if d else None)(last_success_at(session, channel.code)),
            "open_rejections": rejected, "top_errors": [{"message": m, "count": c, "kind": kinds[m]} for m, c in errs.most_common(10)]}


# ---------------------------------------------------------------------------------------------------
# Schedule (EXP-06)
# ---------------------------------------------------------------------------------------------------
DEFAULT_SCHEDULE = {"enabled": False, "mode": "delta", "every_hours": None, "at_hour": None, "at_minute": 0,
                    "timezone": "America/New_York", "view_id": None, "last_run_at": None, "last_result": None}


def schedule_out(channel: Channel) -> dict:
    return {**DEFAULT_SCHEDULE, **(channel.schedule or {})}


def validate_schedule(sched: dict) -> str | None:
    """Returns the reason a schedule is invalid, or None. Used by the API and defensively by the job."""
    if sched.get("mode") not in (None, "delta", "full"):
        return "mode must be delta or full"
    eh, ah, am = sched.get("every_hours"), sched.get("at_hour"), sched.get("at_minute")
    if eh is not None and (not isinstance(eh, (int, float)) or isinstance(eh, bool) or eh <= 0):
        return "every_hours must be greater than zero"
    if ah is not None and (not isinstance(ah, int) or isinstance(ah, bool) or not 0 <= ah <= 23):
        return "at_hour must be 0 to 23"
    if am is not None and (not isinstance(am, int) or isinstance(am, bool) or not 0 <= am <= 59):
        return "at_minute must be 0 to 59"
    if sched.get("enabled") and not eh and ah is None:
        return "Set every_hours or at_hour to enable the schedule"
    return None


def schedule_due(sched: dict, now: datetime | None = None) -> bool:
    """A malformed schedule is never due (and never raises), so one bad channel cannot abort the job."""
    if not sched.get("enabled") or validate_schedule(sched):
        return False
    now = now or utcnow()
    try:
        last = sched.get("last_run_at")
        last_dt = datetime.fromisoformat(last) if last else None
        if sched.get("every_hours"):
            return last_dt is None or now - last_dt >= timedelta(hours=float(sched["every_hours"]))
        if sched.get("at_hour") is not None:
            try:
                from zoneinfo import ZoneInfo

                local = now.astimezone(ZoneInfo(sched.get("timezone") or "America/New_York"))
            except Exception:  # noqa: BLE001
                local = now
            due_today = local.replace(hour=int(sched["at_hour"]), minute=int(sched.get("at_minute") or 0), second=0, microsecond=0)
            if local < due_today:
                return False
            return last_dt is None or last_dt.astimezone(local.tzinfo) < due_today
    except (TypeError, ValueError) as exc:
        log.warning("schedule ignored, invalid values: %s", exc)
    return False


def run_scheduled(session: Session, meta: Metadata, channel: Channel, requested_by: str = "scheduler", force: bool = False) -> dict:
    """One scheduled export; skipped (not queued) when a run is already in progress for the channel."""
    sched = schedule_out(channel)
    if running_run(session, channel.code):
        channel.schedule = {**sched, "last_result": "skipped: a run is already in progress"}
        return {"status": "skipped", "reason": "A run for this channel is already in progress"}
    if not force and not schedule_due(sched):
        return {"status": "skipped", "reason": "Not due"}
    selection = {"view_id": sched["view_id"]} if sched.get("view_id") else {"all": True}
    rec = start_run(session, channel, "scheduled", sched.get("mode", "delta"), requested_by, selection)
    session.flush()
    if selection.get("all"):
        rec.selection = {"entity_ids": [e.id for e in session.execute(select(ProductEntity).where(ProductEntity.level == "style_color")).scalars()]}
    execute_run(session, meta, channel, rec, send=True)
    channel.schedule = {**sched, "last_run_at": utcnow().isoformat(),
                        "last_result": f"{rec.status}: {rec.valid_count} valid, {rec.invalid_count} invalid" + (f" ({rec.error})" if rec.error else "")}
    return {"status": "started", "publish_id": rec.id, "result": rec.status}


def now_utc() -> datetime:
    return datetime.now(UTC)

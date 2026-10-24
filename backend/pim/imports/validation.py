"""IMP-04 validation and IMP-03 value mapping for imports. Every staged row is checked against attribute
definitions, option lists (through remembered per-source value mappings), hierarchy rules, and required flags.
Nothing is written to product data here."""

from __future__ import annotations

import csv
import io
from collections import Counter
from typing import Any

from openpyxl import Workbook
from rapidfuzz import fuzz, process
from sqlalchemy import select
from sqlalchemy.orm import Session

from pim.imports.transforms import map_row
from pim.models import ImportBatch, ValueMapping
from pim.normalization import normalize_raw_option
from pim.services.metadata import LEVEL_ORDER, Metadata
from pim.services.values import NA_SENTINEL, ValueError_, coerce

CORE_KEYS = {"style_number", "color_code", "size", "sku", "upc", "family_code"}


def value_mappings_for(session: Session, source_code: str, list_code: str) -> dict[str, str | None]:
    rows = session.execute(select(ValueMapping).where(ValueMapping.source_code == source_code, ValueMapping.list_code == list_code)).scalars()
    return {r.raw_normalized: r.option_code for r in rows}


def validate_batch(session: Session, meta: Metadata, batch: ImportBatch) -> dict[str, Any]:
    on_error = (batch.options or {}).get("on_error", "skip_value")
    required = [m.get("target") for m in (batch.mapping or []) if m.get("required") and m.get("target")]
    norm = meta.normalizer
    placeholders: Counter = Counter()
    stats = {"rows": 0, "ok": 0, "warnings": 0, "errors": 0}
    vm_cache: dict[str, dict[str, str | None]] = {}
    for row in batch.rows:
        stats["rows"] += 1
        row.errors, row.warnings = [], []
        mapped = map_row(row.raw, batch.mapping or [])
        if row.mapped.get("_pim_id"):
            mapped["_pim_id"] = row.mapped["_pim_id"]
        clean: dict[str, Any] = {}
        for target, raw_val in mapped.items():
            if target in CORE_KEYS or target.startswith("_"):
                clean[target] = str(raw_val).strip() if raw_val not in (None, "") else None
                continue
            attr = meta.attributes.get(target)
            if attr is None:
                row.warnings.append({"attribute": target, "message": "Unknown attribute; ignored"})
                continue
            if attr.level != batch.level and LEVEL_ORDER.get(attr.level, 9) > LEVEL_ORDER.get(batch.level, 0):
                row.errors.append({"attribute": target, "message": f"{attr.label} lives at the {attr.level} level and cannot be set from a {batch.level} row"})
                continue
            if attr.level != batch.level:
                row.warnings.append({"attribute": target, "message": f"{attr.label} lives at the {attr.level} level; it will be written there"})
            if attr.is_erp or (meta.groups.get(attr.group_code) and meta.groups[attr.group_code].read_only):
                row.errors.append({"attribute": target, "message": f"{attr.label} is ERP-owned; imports cannot set it (P1)"})
                continue
            n = norm.normalize_text(raw_val, allows_not_applicable=attr.allows_not_applicable, not_required_means_na=attr.not_required_means_na)
            if n.was_placeholder:
                placeholders[target] += 1
            value: Any = NA_SENTINEL if n.not_applicable else n.value
            if value is not None and value != NA_SENTINEL and attr.data_type in ("single_select", "multi_select") and attr.option_list_code:
                vm = vm_cache.setdefault(attr.option_list_code, value_mappings_for(session, batch.source_code, attr.option_list_code))
                codes = meta.option_codes(attr.option_list_code)
                parts = [p.strip() for p in str(value).split(",")] if attr.data_type == "multi_select" else [str(value)]
                resolved = []
                for part in parts:
                    key = normalize_raw_option(part)
                    code = vm.get(key)
                    if code is None and part in codes:
                        code = part
                    if code is None:
                        label_hit = next((v.code for v in meta.option_lists[attr.option_list_code].values if normalize_raw_option(v.label) == key), None)
                        code = label_hit
                    if code is None:
                        row.errors.append({"attribute": target, "message": f"'{part}' has no value mapping to {attr.label}; map it in the Map step"})
                        break
                    resolved.append(code)
                else:
                    value = resolved if attr.data_type == "multi_select" else resolved[0]
                if len(resolved) != len(parts):
                    continue
            if value is None:
                clean[target] = None
                continue
            try:
                coerce(meta, attr, value)
                clean[target] = value
            except ValueError_ as exc:
                row.errors.append({"attribute": target, "message": str(exc)})
        # hierarchy rules
        if batch.level in ("style_color", "sku") and not clean.get("_pim_id"):
            has_key = any(clean.get(k) for k in ("upc", "sku")) or (clean.get("style_number") and clean.get("color_code"))
            if not has_key:
                row.warnings.append({"message": "No identifier (UPC, SKU, or style and color); assisted matching will be needed. Keyed worksheets avoid this step."})
        for req in required:
            if clean.get(req) in (None, ""):
                row.errors.append({"attribute": req, "message": f"{meta.attributes[req].label if req in meta.attributes else req} is required by this import"})
        if row.errors and on_error == "skip_row":
            row.action = "skip"
        elif row.errors:
            bad = {e.get("attribute") for e in row.errors if e.get("attribute")}
            clean = {k: v for k, v in clean.items() if k not in bad}
            row.action = "update" if row.action == "skip" and not row.decided_by else row.action
        row.mapped = clean
        if row.errors:
            stats["errors"] += 1
        elif row.warnings:
            stats["warnings"] += 1
        else:
            stats["ok"] += 1
    if batch.worksheet_id:
        from pim.imports.worksheet import verify_worksheet_rows
        from pim.models import WorksheetExport

        ws_export = session.get(WorksheetExport, batch.worksheet_id)
        if ws_export is not None:
            rejected = verify_worksheet_rows(session, batch.rows, ws_export)
            stats["errors"] = sum(1 for r in batch.rows if r.errors)
            stats["ok"] = sum(1 for r in batch.rows if not r.errors and not r.warnings)
            stats["worksheet_rows_rejected"] = rejected
    batch.stats = {**(batch.stats or {}), **stats, "placeholders_removed": dict(placeholders),
                   "errors_by_type": dict(Counter(e["message"].split(";")[0][:60] for r in batch.rows for e in r.errors))}
    batch.status = "validated"
    return batch.stats


def value_mapping_report(session: Session, meta: Metadata, batch: ImportBatch) -> list[dict]:
    """IMP-03: distinct incoming values per select attribute with the remembered mapping and a suggestion."""
    out = []
    for m in batch.mapping or []:
        attr = meta.attributes.get(m.get("target") or "")
        if attr is None or attr.data_type not in ("single_select", "multi_select") or not attr.option_list_code:
            continue
        counts: Counter = Counter()
        for row in batch.rows:
            v = map_row(row.raw, [m]).get(attr.code)
            if v:
                for part in (str(v).split(",") if attr.data_type == "multi_select" else [str(v)]):
                    if part.strip():
                        counts[part.strip()] += 1
        existing = value_mappings_for(session, batch.source_code, attr.option_list_code)
        options = {v.code: v.label for v in meta.option_lists[attr.option_list_code].values}
        values = []
        for raw, n in counts.most_common():
            key = normalize_raw_option(raw)
            current = existing.get(key)
            suggested, conf = None, None
            if current is None and options:
                hit = process.extractOne(raw, options, scorer=fuzz.token_set_ratio, processor=lambda s: str(s).lower())
                if hit and hit[1] >= 70:
                    suggested, conf = hit[2], round(hit[1] / 100, 2)
            values.append({"raw": raw, "count": n, "option_code": current, "suggested": suggested, "confidence": conf})
        out.append({"attribute_code": attr.code, "list_code": attr.option_list_code, "values": values})
    return out


def save_value_mappings(session: Session, source_code: str, items: list[dict], user_email: str) -> int:
    n = 0
    for it in items:
        key = normalize_raw_option(it["raw"])
        existing = session.execute(select(ValueMapping).where(ValueMapping.source_code == source_code,
                                                              ValueMapping.list_code == it["list_code"],
                                                              ValueMapping.raw_normalized == key)).scalar_one_or_none()
        if existing is None:
            session.add(ValueMapping(source_code=source_code, list_code=it["list_code"], raw_value=str(it["raw"]), raw_normalized=key,
                                     option_code=it.get("option_code"), created_by=user_email))
        else:
            existing.option_code = it.get("option_code")
        n += 1
    return n


def error_file(batch: ImportBatch, fmt: str = "csv") -> tuple[str, bytes]:
    """IMP-04: the source layout with an added Errors column, so the sender can fix and resend."""
    files = batch.files or []
    rows_by_file: dict[int, list] = {}
    for r in batch.rows:
        rows_by_file.setdefault(r.file_index, []).append(r)
    if fmt == "xlsx":
        wb = Workbook()
        wb.remove(wb.active)
        for i, f in enumerate(files):
            ws = wb.create_sheet((f.get("sheet") or f.get("name") or f"file{i}")[:28])
            headers = [c["name"] for c in f.get("columns", [])] + ["Errors"]
            ws.append(headers)
            for r in rows_by_file.get(i, []):
                ws.append([r.raw.get(h, "") for h in headers[:-1]] + ["; ".join(e["message"] for e in r.errors)])
        buf = io.BytesIO()
        wb.save(buf)
        return "import_errors.xlsx", buf.getvalue()
    f = files[0] if files else {"columns": []}
    headers = [c["name"] for c in f.get("columns", [])] + ["Errors"]
    out = io.StringIO()
    w = csv.DictWriter(out, fieldnames=headers, extrasaction="ignore")
    w.writeheader()
    for r in rows_by_file.get(0, []):
        w.writerow({**{h: r.raw.get(h, "") for h in headers[:-1]}, "Errors": "; ".join(e["message"] for e in r.errors)})
    return "import_errors.csv", out.getvalue().encode("utf-8")

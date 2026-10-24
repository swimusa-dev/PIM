"""IMP-12 keyed worksheets: an XLSX laid out like the team's setup file (group header row, field row, frozen panes)
with hidden, locked key columns (internal id, style number, color code, checksum), dropdowns from option lists,
shaded read-only ERP context columns, and an Instructions tab. Returned files are recognized by the embedded id."""

from __future__ import annotations

import hashlib
import io
import secrets
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Protection
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from sqlalchemy.orm import Session

from pim.auth.identity import CurrentUser
from pim.imports.parsers import WORKSHEET_SHEET
from pim.models import ProductEntity, WorksheetExport
from pim.services.metadata import Metadata
from pim.services.values import resolve_cells

TEAL = "3B859C"
KEY_COLUMNS = ["_pim_id", "_style_number", "_color_code", "_checksum"]


def checksum(secret: str, entity_id: str, style: str, color: str | None) -> str:
    return hashlib.sha256(f"{secret}|{entity_id}|{style}|{color or ''}".encode()).hexdigest()[:12]


def build_worksheet(session: Session, meta: Metadata, user: CurrentUser, entities: list[ProductEntity], level: str,
                    group_codes: list[str], include_erp_context: bool, selection: dict) -> tuple[WorksheetExport, bytes]:
    secret = secrets.token_hex(8)
    groups = [meta.groups[g] for g in group_codes if g in meta.groups and g != "erp"]
    attrs = [a for g in groups for a in sorted(g.attributes, key=lambda a: (a.sort_order, a.code))
             if a.level == level and a.status != "deprecated" and not a.is_erp and a.data_type != "reference_entity"]
    erp_attrs = [a for a in meta.attributes_for_level(level) if a.is_erp][:8] if include_erp_context else []
    codes = [a.code for a in erp_attrs] + [a.code for a in attrs]
    cells = resolve_cells(session, meta, user, entities, codes, include_drafts=False)
    ws_export = WorksheetExport(created_by=user.email, selection=selection, level=level, group_codes=group_codes,
                                attribute_codes=[a.code for a in attrs], secret=secret, row_count=len(entities), file_name="", path="")
    session.add(ws_export)
    session.flush()

    wb = Workbook()
    ws = wb.active
    ws.title = "Product Information"
    headers_group = [*["Key"] * 4, "Product", "Product", *["ERP (read-only)"] * len(erp_attrs)]
    headers_field = [*KEY_COLUMNS, "Style number", "Color code", *[a.label for a in erp_attrs]]
    for g in groups:
        for a in attrs:
            if a.group_code == g.code:
                headers_group.append(g.label)
                headers_field.append(a.label)
    ws.append(headers_group)
    ws.append(headers_field)
    # merge group cells
    start = 1
    for i in range(2, len(headers_group) + 2):
        if i == len(headers_group) + 1 or headers_group[i - 1] != headers_group[start - 1]:
            if i - 1 > start:
                ws.merge_cells(start_row=1, start_column=start, end_row=1, end_column=i - 1)
            start = i
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=TEAL)
    for c in ws[2]:
        c.font = Font(bold=True)
    erp_fill = PatternFill("solid", fgColor="E7E6E6")
    attr_by_label = {a.label: a for a in attrs}
    for ent in entities:
        row_cells = cells.get(ent.id, {})
        row = [ent.id, ent.style_number, ent.color_code or "", checksum(secret, ent.id, ent.style_number, ent.color_code),
               ent.style_number, ent.color_code or ""]
        for a in erp_attrs:
            row.append(_display(row_cells.get(a.code)))
        for h in headers_field[6 + len(erp_attrs):]:
            row.append(_display(row_cells.get(attr_by_label[h].code)))
        ws.append(row)
    n_rows = len(entities) + 2
    # lock everything, then unlock editable attribute cells
    for r in ws.iter_rows(min_row=1, max_row=n_rows):
        for c in r:
            c.protection = Protection(locked=True)
    first_edit = 7 + len(erp_attrs)
    for r in ws.iter_rows(min_row=3, max_row=n_rows, min_col=first_edit, max_col=len(headers_field)):
        for c in r:
            c.protection = Protection(locked=False)
    for r in ws.iter_rows(min_row=3, max_row=n_rows, min_col=7, max_col=6 + len(erp_attrs)):
        for c in r:
            c.fill = erp_fill
    for i in range(1, 5):
        ws.column_dimensions[get_column_letter(i)].hidden = True
    # dropdowns from option lists
    for idx, h in enumerate(headers_field, start=1):
        a = attr_by_label.get(h)
        if a and a.data_type == "single_select" and a.option_list_code in meta.option_lists:
            labels = [v.label for v in meta.option_lists[a.option_list_code].values]
            formula = '"' + ",".join(x.replace('"', "'")[:50] for x in labels)[:250] + '"'
            dv = DataValidation(type="list", formula1=formula, allow_blank=True)
            ws.add_data_validation(dv)
            dv.add(f"{get_column_letter(idx)}3:{get_column_letter(idx)}{max(n_rows, 3)}")
    ws.freeze_panes = "G3"
    ws.protection.sheet = True
    ws.protection.password = secret[:8]
    meta_ws = wb.create_sheet(WORKSHEET_SHEET)
    meta_ws.sheet_state = "hidden"
    meta_ws.append(["worksheet_id", ws_export.id])
    meta_ws.append(["level", level])
    meta_ws.append(["header_row", 2])
    meta_ws.append(["group_header_row", 1])
    for h in headers_field:
        a = attr_by_label.get(h)
        if a:
            meta_ws.append([f"col:{h}", a.code])
    for h, code in (("Style number", "style_number"), ("Color code", "color_code")):
        meta_ws.append([f"col:{h}", code])
    instr = wb.create_sheet("Instructions")
    for line in [
        "Swim USA PIM worksheet",
        "Fill in the white cells only. Gray cells are ERP context and are read-only. Hidden columns A to D identify each row; do not unhide or edit them.",
        "Use the dropdowns where present. Leave a cell blank to make no change; type N/A only where the attribute allows Not applicable.",
        "Dates: YYYY-MM-DD. Measurements: numbers with the unit named in the column header. Do not use the em dash character.",
        "Return the file through Imports with source 'Keyed worksheet'. Rows match by key with no mapping step; changes go through validation, survivorship, and review.",
    ]:
        instr.append([line])
    buf = io.BytesIO()
    wb.save(buf)
    data = buf.getvalue()
    ws_export.file_name = f"pim_worksheet_{ws_export.id[:8]}.xlsx"
    return ws_export, data


def _display(cell: dict | None) -> Any:
    if not cell:
        return ""
    if cell.get("not_applicable"):
        return "N/A"
    v = cell.get("display") or cell.get("value")
    if isinstance(v, list):
        return ", ".join(str(x) for x in v)
    return "" if v is None else v


def worksheet_mapping(meta: Metadata, worksheet_meta: dict, columns: list[str]) -> tuple[list[dict], list[str]]:
    """Mapping for a returned worksheet: field labels -> attribute codes from the hidden sheet; keys are the hidden columns."""
    mapping = [{"target": "_pim_id", "sources": ["_pim_id"], "transforms": []},
               {"target": "_checksum", "sources": ["_checksum"], "transforms": []},
               {"target": "_style_number", "sources": ["_style_number"], "transforms": []},
               {"target": "_color_code", "sources": ["_color_code"], "transforms": []}]
    for k, code in worksheet_meta.items():
        if str(k).startswith("col:") and str(k)[4:] in columns and code not in ("style_number", "color_code"):
            mapping.append({"target": str(code), "sources": [str(k)[4:]], "transforms": [{"op": "trim", "args": {}}]})
    return mapping, ["_pim_id"]


def verify_worksheet_rows(session: Session, batch_rows, ws_export: WorksheetExport) -> int:  # noqa: ANN001
    """Reject edits to locked columns (IMP-12): the checksum must match the hidden key columns."""
    rejected = 0
    for row in batch_rows:
        pid = (row.raw.get("_pim_id") or "").strip()
        style, color = (row.raw.get("_style_number") or "").strip(), (row.raw.get("_color_code") or "").strip() or None
        expected = checksum(ws_export.secret, pid, style, color) if pid else ""
        if not pid or (row.raw.get("_checksum") or "").strip() != expected:
            row.errors.append({"message": "Locked key columns were edited or removed; row rejected"})
            row.action, rejected = "skip", rejected + 1
            continue
        ent = session.get(ProductEntity, pid)
        if ent is None:
            row.errors.append({"message": "Row key no longer exists in the PIM"})
            row.action, rejected = "skip", rejected + 1
            continue
        row.mapped["_pim_id"] = pid
        row.entity_id, row.match_status, row.match_key_used, row.action = pid, "matched", "worksheet", "update"
    return rejected
